from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import CellRange


CELL_ADDRESS_RE = re.compile(r"^([A-Za-z]+)([1-9][0-9]*)$")
COLUMN_RE = re.compile(r"^[A-Za-z]+$")


@dataclass(frozen=True, slots=True)
class Cell:
    """A cell-shaped object returned by the Excel read protocol.

    `range` is the identity. For merged cells it is the full merged A1 range,
    and for ordinary cells it is the single A1 address.
    """

    range: str
    value: Any
    is_merged: bool
    rowspan: int = 1
    colspan: int = 1

    def to_dict(self) -> dict[str, Any]:
        return {
            "range": self.range,
            "value": self.value,
            "is_merged": self.is_merged,
            "rowspan": self.rowspan,
            "colspan": self.colspan,
        }

    def to_compact_dict(self) -> dict[str, Any]:
        payload: dict[str, Any] = {"r": self.range}
        if self.value is not None:
            payload["v"] = _json_safe_value(self.value)
        if self.rowspan != 1:
            payload["rs"] = self.rowspan
        if self.colspan != 1:
            payload["cs"] = self.colspan
        return payload


class ExcelReadProtocol:
    """Controlled, bounded Excel read protocol.

    The core protocol methods are `get_rows`, `get_columns`, and `get_cell`.
    Sheet context is managed outside those methods through `list_sheets`,
    `set_current_sheet`, and `get_current_sheet`.
    """

    def __init__(self, workbook_path: str | Path, *, sheet: str | None = None) -> None:
        self.workbook_path = Path(workbook_path)
        if not self.workbook_path.exists():
            raise FileNotFoundError(self.workbook_path)
        self._workbook = load_workbook(self.workbook_path, data_only=True)
        self._merge_indexes: dict[str, MergeIndex] = {}
        initial_sheet = sheet or self._workbook.active.title
        self.set_current_sheet(initial_sheet)

    def list_sheets(self) -> list[str]:
        return list(self._workbook.sheetnames)

    def set_current_sheet(self, name: str) -> None:
        if name not in self._workbook.sheetnames:
            raise KeyError(f"sheet {name!r} does not exist. Available sheets: {self._workbook.sheetnames}")
        self._current_sheet = name

    def get_current_sheet(self) -> str:
        return self._current_sheet

    def get_rows(self, start: int, end: int) -> list[Cell]:
        """Return cells whose row ownership belongs to rows start..end."""
        sheet = self._sheet()
        start, end = self._validate_row_scan(start, end)
        if start > sheet.max_row:
            return []
        end = min(end, sheet.max_row)

        index = self._merge_index(sheet.title)
        result: list[Cell] = []
        seen_merges: set[str] = set()
        for row in range(start, end + 1):
            for col in range(1, sheet.max_column + 1):
                merged = index.find(row, col)
                if merged is not None:
                    if merged.min_row != row:
                        continue
                    cell = self._build_cell(sheet, index, merged.min_row, merged.min_col)
                    if cell.range in seen_merges:
                        continue
                    seen_merges.add(cell.range)
                    result.append(cell)
                    continue
                result.append(self._build_cell(sheet, index, row, col))
        return result

    def get_columns(self, start: str, end: str) -> list[Cell]:
        """Return cells whose column ownership belongs to columns start..end."""
        sheet = self._sheet()
        start_col, end_col = self._validate_column_scan(start, end)
        if start_col > sheet.max_column:
            return []
        end_col = min(end_col, sheet.max_column)

        index = self._merge_index(sheet.title)
        result: list[Cell] = []
        seen_merges: set[str] = set()
        for col in range(start_col, end_col + 1):
            for row in range(1, sheet.max_row + 1):
                merged = index.find(row, col)
                if merged is not None:
                    if merged.min_col != col:
                        continue
                    cell = self._build_cell(sheet, index, merged.min_row, merged.min_col)
                    if cell.range in seen_merges:
                        continue
                    seen_merges.add(cell.range)
                    result.append(cell)
                    continue
                result.append(self._build_cell(sheet, index, row, col))
        return result

    def get_cell(self, address: str) -> Cell:
        """Return the ordinary cell or full merged block addressed by address."""
        sheet = self._sheet()
        row, col = self._parse_cell_address(address)
        if row > sheet.max_row or col > sheet.max_column:
            raise ValueError(
                f"cell address {address!r} is outside current sheet bounds. "
                f"Valid range for current sheet {sheet.title!r}: A1..{get_column_letter(sheet.max_column)}{sheet.max_row}"
            )
        index = self._merge_index(sheet.title)
        return self._build_cell(sheet, index, row, col)

    def _sheet(self) -> Any:
        return self._workbook[self._current_sheet]

    def _merge_index(self, sheet_name: str) -> "MergeIndex":
        if sheet_name not in self._merge_indexes:
            self._merge_indexes[sheet_name] = MergeIndex(self._workbook[sheet_name])
        return self._merge_indexes[sheet_name]

    def _build_cell(self, sheet: Any, merge_index: "MergeIndex", row: int, col: int) -> Cell:
        merged = merge_index.find(row, col)
        if merged is not None:
            value = sheet.cell(row=merged.min_row, column=merged.min_col).value
            return Cell(
                range=_range_ref(merged),
                value=value,
                is_merged=True,
                rowspan=merged.max_row - merged.min_row + 1,
                colspan=merged.max_col - merged.min_col + 1,
            )
        return Cell(
            range=f"{get_column_letter(col)}{row}",
            value=sheet.cell(row=row, column=col).value,
            is_merged=False,
            rowspan=1,
            colspan=1,
        )

    def _validate_row_scan(self, start: int, end: int) -> tuple[int, int]:
        sheet = self._sheet()
        try:
            start_int = int(start)
            end_int = int(end)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"row range {start!r}..{end!r} is invalid; expected 1-based integers. "
                f"Valid row range for current sheet {sheet.title!r}: 1..{sheet.max_row}"
            ) from exc
        if start_int < 1:
            raise ValueError(
                f"row {start!r} is invalid (rows are 1-based). "
                f"Valid range for current sheet {sheet.title!r}: 1..{sheet.max_row}"
            )
        if start_int > end_int:
            raise ValueError(
                f"row range {start_int}..{end_int} is invalid; start must be <= end. "
                f"Valid range for current sheet {sheet.title!r}: 1..{sheet.max_row}"
            )
        return start_int, end_int

    def _validate_column_scan(self, start: str, end: str) -> tuple[int, int]:
        sheet = self._sheet()
        start_col = self._parse_column(start, "start")
        end_col = self._parse_column(end, "end")
        if start_col > end_col:
            raise ValueError(
                f"column range {start!r}..{end!r} is invalid; start must be <= end. "
                f"Valid column range for current sheet {sheet.title!r}: A..{get_column_letter(sheet.max_column)}"
            )
        return start_col, end_col

    def _parse_column(self, value: str, label: str) -> int:
        sheet = self._sheet()
        text = str(value or "").strip().upper()
        if not COLUMN_RE.fullmatch(text):
            raise ValueError(
                f"{label} column {value!r} is invalid; expected column letters like 'A' or 'AA'. "
                f"Valid column range for current sheet {sheet.title!r}: A..{get_column_letter(sheet.max_column)}"
            )
        return column_index_from_string(text)

    def _parse_cell_address(self, address: str) -> tuple[int, int]:
        sheet = self._sheet()
        text = str(address or "").strip().upper()
        if ":" in text:
            raise ValueError(
                f"cell address {address!r} is invalid; get_cell expects one A1 address, not a range. "
                f"Valid range for current sheet {sheet.title!r}: A1..{get_column_letter(sheet.max_column)}{sheet.max_row}"
            )
        match = CELL_ADDRESS_RE.fullmatch(text)
        if match is None:
            raise ValueError(
                f"cell address {address!r} is invalid; expected A1 notation like 'B5'. "
                f"Valid range for current sheet {sheet.title!r}: A1..{get_column_letter(sheet.max_column)}{sheet.max_row}"
            )
        col_text, row_text = match.groups()
        return int(row_text), column_index_from_string(col_text)


class MergeIndex:
    def __init__(self, sheet: Any) -> None:
        self._map: dict[tuple[int, int], CellRange] = {}
        for merged in sheet.merged_cells.ranges:
            for row in range(merged.min_row, merged.max_row + 1):
                for col in range(merged.min_col, merged.max_col + 1):
                    self._map[(row, col)] = merged

    def find(self, row: int, col: int) -> CellRange | None:
        return self._map.get((row, col))


def _range_ref(merged: CellRange) -> str:
    return f"{get_column_letter(merged.min_col)}{merged.min_row}:{get_column_letter(merged.max_col)}{merged.max_row}"


def _json_safe_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value

