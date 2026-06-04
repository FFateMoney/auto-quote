from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.cell_range import CellRange


RANGE_RE = re.compile(r"^[A-Za-z]+[1-9][0-9]*(?::[A-Za-z]+[1-9][0-9]*)?$")


@dataclass(frozen=True, slots=True)
class Chunk:
    label: str | None
    main_range: str
    shared_ranges: list[str]
    workbook: Workbook


class ExcelChunkProtocol:
    """Create independent workbook chunks from the current sheet.

    The public protocol is intentionally one function: `create_chunk`.
    Sheet context is managed through `list_sheets`, `set_current_sheet`, and
    `get_current_sheet`, matching the read protocol.
    """

    def __init__(self, workbook_path: str | Path, *, sheet: str | None = None) -> None:
        self.workbook_path = Path(workbook_path)
        if not self.workbook_path.exists():
            raise FileNotFoundError(self.workbook_path)
        self._workbook = load_workbook(self.workbook_path, data_only=True)
        initial_sheet = sheet or self._workbook.active.title
        self.set_current_sheet(initial_sheet)

    def list_sheets(self) -> list[str]:
        return list(self._workbook.sheetnames)

    def set_current_sheet(self, name: str) -> None:
        if name not in self._workbook.sheetnames:
            raise KeyError(f"sheet {name!r} does not exist. Available sheets: {self._workbook.sheetnames}")
        self._current_sheet = name

    def get_current_sheet(self) -> str:
        return self._current_sheet

    def create_chunk(
        self,
        main_range: str,
        shared_ranges: list[str] | None = None,
        label: str | None = None,
    ) -> Chunk:
        source_sheet = self._sheet()
        normalized_main = self._normalize_required_range("main_range", main_range)
        normalized_shared = [
            self._normalize_required_range(f"shared_ranges[{index}]", item)
            for index, item in enumerate(shared_ranges or [])
        ]

        all_ranges = [("main_range", normalized_main), *[
            (f"shared_ranges[{index}]", item) for index, item in enumerate(normalized_shared)
        ]]
        for name, ref in all_ranges:
            self._validate_bounds(name, ref, source_sheet)
        for name, ref in all_ranges:
            self._validate_merged_cells(name, ref, source_sheet)

        workbook = self._materialize(source_sheet, [normalized_main, *normalized_shared])
        return Chunk(label=label, main_range=normalized_main, shared_ranges=normalized_shared, workbook=workbook)

    def _sheet(self) -> Any:
        return self._workbook[self._current_sheet]

    def _normalize_required_range(self, name: str, value: str) -> str:
        if value is None:
            raise ValueError(f'{name} is required and must be a non-empty A1 range, got None.')
        text = str(value).strip().upper()
        if not text:
            raise ValueError(f'{name} is required and must be a non-empty A1 range, got {value!r}.')
        if not RANGE_RE.fullmatch(text):
            raise ValueError(
                f'{name} "{value}" is not a valid A1 range. Expected forms like "B2:D7" or "B5".'
            )
        try:
            min_col, min_row, max_col, max_row = range_boundaries(text)
        except ValueError as exc:
            raise ValueError(
                f'{name} "{value}" is not a valid A1 range. Expected forms like "B2:D7" or "B5".'
            ) from exc
        if min_col > max_col or min_row > max_row:
            suggestion = _format_range(max_col, max_row, min_col, min_row)
            raise ValueError(
                f'{name} "{value}" is not a valid range '
                f"(top-left corner must come before bottom-right corner). "
                f'Did you mean "{suggestion}"?'
            )
        return _format_range(min_col, min_row, max_col, max_row)

    def _validate_bounds(self, name: str, ref: str, sheet: Any) -> None:
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        if min_row < 1 or min_col < 1 or max_row > sheet.max_row or max_col > sheet.max_column:
            raise ValueError(
                f'{name} "{ref}" exceeds sheet bounds. '
                f"Current sheet max_row={sheet.max_row}, max_col={sheet.max_column} "
                f"(A-{get_column_letter(sheet.max_column)}). "
                f"Valid range: A1:{get_column_letter(sheet.max_column)}{sheet.max_row}."
            )

    def _validate_merged_cells(self, name: str, ref: str, sheet: Any) -> None:
        rng = range_boundaries(ref)
        for merged in _sorted_merged_ranges(sheet):
            merge = _merged_bounds(merged)
            if _cuts_through(merge, rng):
                raise ValueError(_merge_cut_error(name, ref, merge, rng))

    def _materialize(self, source_sheet: Any, refs: list[str]) -> Workbook:
        cells_to_copy: set[tuple[int, int]] = set()
        for ref in refs:
            min_col, min_row, max_col, max_row = range_boundaries(ref)
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    cells_to_copy.add((row, col))

        workbook = Workbook()
        workbook.remove(workbook.active)
        target_sheet = workbook.create_sheet(title=source_sheet.title)

        for row, col in sorted(cells_to_copy):
            value = source_sheet.cell(row=row, column=col).value
            if value is not None:
                target_sheet.cell(row=row, column=col, value=value)

        for merged in _sorted_merged_ranges(source_sheet):
            merge_cells = [
                (row, col)
                for row in range(merged.min_row, merged.max_row + 1)
                for col in range(merged.min_col, merged.max_col + 1)
            ]
            if all(cell in cells_to_copy for cell in merge_cells):
                target_sheet.merge_cells(str(merged))

        return workbook


def _merged_bounds(merged: CellRange) -> tuple[int, int, int, int]:
    return merged.min_col, merged.min_row, merged.max_col, merged.max_row


def _sorted_merged_ranges(sheet: Any) -> list[CellRange]:
    return sorted(sheet.merged_cells.ranges, key=lambda item: (item.min_row, item.min_col, item.max_row, item.max_col))


def _intersects(left: tuple[int, int, int, int], right: tuple[int, int, int, int]) -> bool:
    left_min_col, left_min_row, left_max_col, left_max_row = left
    right_min_col, right_min_row, right_max_col, right_max_row = right
    return (
        left_min_row <= right_max_row
        and left_max_row >= right_min_row
        and left_min_col <= right_max_col
        and left_max_col >= right_min_col
    )


def _is_fully_inside(inner: tuple[int, int, int, int], outer: tuple[int, int, int, int]) -> bool:
    inner_min_col, inner_min_row, inner_max_col, inner_max_row = inner
    outer_min_col, outer_min_row, outer_max_col, outer_max_row = outer
    return (
        outer_min_row <= inner_min_row
        and inner_max_row <= outer_max_row
        and outer_min_col <= inner_min_col
        and inner_max_col <= outer_max_col
    )


def _cuts_through(merge: tuple[int, int, int, int], rng: tuple[int, int, int, int]) -> bool:
    return _intersects(merge, rng) and not _is_fully_inside(merge, rng)


def _format_range(min_col: int, min_row: int, max_col: int, max_row: int) -> str:
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def _merge_cut_error(
    name: str,
    ref: str,
    merge: tuple[int, int, int, int],
    rng: tuple[int, int, int, int],
) -> str:
    merge_ref = _format_range(*merge)
    extended = _format_range(
        min(rng[0], merge[0]),
        min(rng[1], merge[1]),
        max(rng[2], merge[2]),
        max(rng[3], merge[3]),
    )
    detail = _merge_cut_detail(merge, rng)
    excluded = _exclude_merge_from_range(rng, merge)
    preferred = (
        "  - Preferred for batch splitting: keep main_range limited to the quote-specific cells, "
        f'and move the full merged cell into shared_ranges: "{merge_ref}"\n'
    )
    if excluded:
        preferred += f'  - If this merged cell is shared context, retry with main_range="{excluded}" and add "{merge_ref}" to shared_ranges.\n'
    return (
        f'{name} "{ref}" cuts through merged cell "{merge_ref}" '
        f"({detail}). "
        "Suggestions:\n"
        f"{preferred}"
        f'  - Only if the merged cell truly belongs to this single quote item, extend the range to include it: "{extended}"'
    )


def _merge_cut_detail(merge: tuple[int, int, int, int], rng: tuple[int, int, int, int]) -> str:
    merge_min_col, merge_min_row, merge_max_col, merge_max_row = merge
    range_min_col, range_min_row, range_max_col, range_max_row = rng
    if merge_max_col > range_max_col:
        return f"merged cell extends to column {get_column_letter(merge_max_col)}, but range ends at column {get_column_letter(range_max_col)}"
    if merge_min_col < range_min_col:
        return f"merged cell starts at column {get_column_letter(merge_min_col)}, but range starts at column {get_column_letter(range_min_col)}"
    if merge_max_row > range_max_row:
        return f"merged cell extends to row {merge_max_row}, but range ends at row {range_max_row}"
    if merge_min_row < range_min_row:
        return f"merged cell starts at row {merge_min_row}, but range starts at row {range_min_row}"
    return "range partially intersects the merged cell"


def _exclude_merge_from_range(
    rng: tuple[int, int, int, int],
    merge: tuple[int, int, int, int],
) -> str:
    range_min_col, range_min_row, range_max_col, range_max_row = rng
    merge_min_col, merge_min_row, merge_max_col, merge_max_row = merge
    if range_min_row < merge_min_row <= range_max_row:
        return _format_range(range_min_col, range_min_row, range_max_col, merge_min_row - 1)
    if range_min_row <= merge_max_row < range_max_row:
        return _format_range(range_min_col, merge_max_row + 1, range_max_col, range_max_row)
    if range_min_col < merge_min_col <= range_max_col:
        return _format_range(range_min_col, range_min_row, merge_min_col - 1, range_max_row)
    if range_min_col <= merge_max_col < range_max_col:
        return _format_range(merge_max_col + 1, range_min_row, range_max_col, range_max_row)
    return ""
