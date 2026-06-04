from __future__ import annotations

import json
import logging
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openai import OpenAI
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from backend.common.logging import append_run_log
from backend.common.models import NormalizedDocument, NormalizedTextBlock
from backend.quote.models import ExtraStandardRequirement, FormRow, StandardContextDecision, StandardEvidence
from backend.quote.settings import get_settings


DOCUMENT_EXTRACT_VISIBLE_FIELDS: tuple[str, ...] = (
    "raw_test_type",
    "canonical_test_type",
    "standard_codes",
    "pricing_mode",
    "pricing_quantity",
    "sample_count",
    "sample_length_mm",
    "sample_width_mm",
    "sample_height_mm",
    "sample_weight_kg",
    "source_text",
    "conditions_text",
    "sample_info_text",
)

DOCUMENT_TARGETED_CONTEXT_FIELDS: tuple[str, ...] = (
    "row_id",
    "raw_test_type",
    "canonical_test_type",
    "standard_codes",
    "planned_standard_fields",
    "source_text",
    "conditions_text",
    "sample_info_text",
)

STANDARD_ENRICH_VISIBLE_FIELDS: tuple[str, ...] = (
    "row_id",
    "raw_test_type",
    "canonical_test_type",
    "standard_codes",
    "pricing_mode",
    "pricing_quantity",
    "sample_length_mm",
    "sample_width_mm",
    "sample_height_mm",
    "sample_weight_kg",
    "required_temp_min",
    "required_temp_max",
    "required_humidity_min",
    "required_humidity_max",
    "required_temp_change_rate",
    "required_freq_min",
    "required_freq_max",
    "required_accel_min",
    "required_accel_max",
    "required_displacement_min",
    "required_displacement_max",
    "required_irradiance_min",
    "required_irradiance_max",
    "required_water_temp_min",
    "required_water_temp_max",
    "required_water_flow_min",
    "required_water_flow_max",
    "planned_standard_fields",
    "discovered_standard_fields",
    "source_text",
    "conditions_text",
    "sample_info_text",
)

STANDARD_DISCOVERY_VISIBLE_FIELDS: tuple[str, ...] = (
    "row_id",
    "raw_test_type",
    "canonical_test_type",
    "standard_codes",
    "planned_standard_fields",
    "required_temp_min",
    "required_temp_max",
    "required_humidity_min",
    "required_humidity_max",
    "required_temp_change_rate",
    "required_freq_min",
    "required_freq_max",
    "required_accel_min",
    "required_accel_max",
    "required_displacement_min",
    "required_displacement_max",
    "required_irradiance_min",
    "required_irradiance_max",
    "required_water_temp_min",
    "required_water_temp_max",
    "required_water_flow_min",
    "required_water_flow_max",
    "source_text",
    "conditions_text",
    "sample_info_text",
)

CONTEXT_JUDGE_VISIBLE_FIELDS: tuple[str, ...] = (
    "row_id",
    "raw_test_type",
    "canonical_test_type",
    "standard_codes",
    "sample_length_mm",
    "sample_width_mm",
    "sample_height_mm",
    "sample_weight_kg",
    "required_temp_min",
    "required_temp_max",
    "required_humidity_min",
    "required_humidity_max",
    "required_temp_change_rate",
    "required_freq_min",
    "required_freq_max",
    "required_accel_min",
    "required_accel_max",
    "required_displacement_min",
    "required_displacement_max",
    "required_irradiance_min",
    "required_irradiance_max",
    "required_water_temp_min",
    "required_water_temp_max",
    "required_water_flow_min",
    "required_water_flow_max",
    "source_text",
    "conditions_text",
    "sample_info_text",
)

CODE_FENCE_RE = re.compile(r"^```(?:json)?\s*|\s*```$", re.DOTALL)
RANGE_FIELD_PAIRS: tuple[tuple[str, str], ...] = (
    ("required_temp_min", "required_temp_max"),
    ("required_humidity_min", "required_humidity_max"),
    ("required_freq_min", "required_freq_max"),
    ("required_accel_min", "required_accel_max"),
    ("required_displacement_min", "required_displacement_max"),
    ("required_irradiance_min", "required_irradiance_max"),
    ("required_water_temp_min", "required_water_temp_max"),
    ("required_water_flow_min", "required_water_flow_max"),
)
logger = logging.getLogger(__name__)


@dataclass(slots=True)
class ModelFillResult:
    items: list[FormRow]
    summary: str = ""
    raw_response: str = ""


@dataclass(slots=True)
class BatchQuoteSplitItem:
    quote_id: str
    title: str
    source_summary: str
    text: str
    block_ids: list[str]


@dataclass(slots=True)
class BatchQuoteSplitResult:
    items: list[BatchQuoteSplitItem]
    summary: str = ""
    raw_response: str = ""


@dataclass(slots=True)
class BatchExcelChunkItem:
    quote_id: str
    title: str
    source_summary: str
    chunk: Any


@dataclass(slots=True)
class BatchExcelChunkSplitResult:
    items: list[BatchExcelChunkItem]
    summary: str = ""
    raw_response: str = ""


@dataclass(slots=True)
class StandardFieldDiscoveryItem:
    row_id: str
    discovered_standard_fields: list[str]
    extra_standard_requirements: list[ExtraStandardRequirement]


@dataclass(slots=True)
class StandardFieldDiscoveryResult:
    items: list[StandardFieldDiscoveryItem]
    summary: str = ""
    raw_response: str = ""


def _load_prompts(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _schema_example(visible_fields: tuple[str, ...]) -> dict[str, Any]:
    payload = {field: _schema_placeholder(field) for field in visible_fields}
    return {"items": [payload]}


def _schema_placeholder(field: str) -> Any:
    if field in {"standard_codes", "planned_standard_fields", "discovered_standard_fields"}:
        return []
    if field == "extra_standard_requirements":
        return [{"requirement_name": "", "requirement_text": "", "source_section": ""}]
    if field in {
        "pricing_quantity", "sample_count", "repeat_count",
        "sample_length_mm", "sample_width_mm", "sample_height_mm", "sample_weight_kg",
        "required_temp_min", "required_temp_max", "required_humidity_min", "required_humidity_max",
        "required_temp_change_rate", "required_freq_min", "required_freq_max",
        "required_accel_min", "required_accel_max", "required_displacement_min", "required_displacement_max",
        "required_irradiance_min", "required_irradiance_max",
        "required_water_temp_min", "required_water_temp_max",
        "required_water_flow_min", "required_water_flow_max",
        "matched_test_type_id", "base_fee", "unit_price", "total_price",
    }:
        return None
    return ""


def _document_targeted_visible_fields(target_fields_by_row: dict[str, list[str]]) -> tuple[str, ...]:
    fields = [*DOCUMENT_TARGETED_CONTEXT_FIELDS, "extra_standard_requirements"]
    for targets in target_fields_by_row.values():
        for field in targets:
            if field not in fields:
                fields.append(field)
    return tuple(fields)


def _test_type_options_text(test_type_options: list[str]) -> str:
    options = [str(item or "").strip() for item in test_type_options if str(item or "").strip()]
    if not options:
        return "- (未加载到标准试验类型列表；canonical_test_type 无法确认时填空字符串)"
    return "\n".join(f"- {name}" for name in options)


def _jsonable_value(value: Any) -> Any:
    if hasattr(value, "model_dump"):
        return value.model_dump()
    if isinstance(value, list):
        return [_jsonable_value(item) for item in value]
    if isinstance(value, dict):
        return {key: _jsonable_value(item) for key, item in value.items()}
    return value


def _rules_text(*, preserve_row_ids: bool, include_sample_count: bool) -> str:
    lines = [
        "只输出 JSON，对象根节点必须是 items 数组。",
        "一个测试项目对应一行，不要把同一项目重复拆成多行。",
        "`standard_codes` 必须是字符串数组，没有就给空数组。",
        "`sample_length_mm`、`sample_width_mm`、`sample_height_mm` 分别填写样品长宽高，未知时给 null；如果文档没有直接写样品长宽，但明确写了夹具或工装的长宽，且该尺寸显然对应待测件装夹/占位尺寸，可将夹具或工装长宽视为样品长宽填写。",
        "数值字段只能填数字或 null，不要带单位。",
        "`source_text`、`conditions_text`、`sample_info_text` 使用简洁中文总结来源信息。",
        "遇到无法确定的字段留空字符串、空数组或 null，不要编造。",
        "如果同一信息存在中文、英文或其他语言版本且含义冲突，以中文内容为准。",
        "如果文档写的是确定单值而不是范围，例如温度 80℃、湿度 95%RH、水温 25℃、流量 10L/min、辐照 800W/m2，就把对应的 min/max 两个字段都填成同一个值。",
        "文中的 `[IMAGE_n]` 与图片输入一一对应，图片是文档插图，不是时间序列视频帧。",
        "`pricing_mode` 表示计价单位或计价方式，例如 小时、件、次、天、组、轴、方向。",
        "`pricing_quantity` 表示本报价行的总计价数量；如果文档写“每方向/每轴 5 小时，共 3 个方向/轴”，应填写 15。",
        "`required_temp_change_rate` 是唯一允许通过计算或推断得到的字段：如果文档或标准明确给出温变速率，则直接填写数值部分；如果没有直接给出，但给出了温度范围和对应升温/降温耗时，可以据此计算并填写；如果只有温度范围没有对应耗时，或只有总时长但无法确认对应的是升降温阶段，就不要猜。",
    ]
    if include_sample_count:
        lines.insert(10, "`sample_count` 表示本次测试的总件数/总样品数；如果同一报价需求包含多个颜色、型号、样品组或件数，应加总；如果文档没有明确写出就留空。")
    if preserve_row_ids:
        lines.append("如果是在补全已有表格，必须尽量保留已有行的 `row_id`，并在原有行上补字段，不要新增重复行。")
    lines = [f"{idx}. {line}" for idx, line in enumerate(lines, start=1)]
    return "\n".join(lines)


def _flatten_stream_part(parts: list[Any]) -> str:
    chunks: list[str] = []
    for part in parts:
        if isinstance(part, str):
            chunks.append(part)
        elif isinstance(part, dict) and isinstance(part.get("text"), str):
            chunks.append(part["text"])
        else:
            text = getattr(part, "text", None)
            if isinstance(text, str):
                chunks.append(text)
    return "".join(chunks)


def _strip_code_fence(text: str) -> str:
    clean = (text or "").strip()
    if clean.startswith("```"):
        clean = CODE_FENCE_RE.sub("", clean).strip()
    return clean


def _extract_json_text(text: str) -> str:
    clean = _strip_code_fence(text)
    match = re.search(r"(\{.*\}|\[.*\])", clean, re.DOTALL)
    return match.group(1) if match else clean


def _normalize_item_payload(item: dict[str, Any]) -> dict[str, Any]:
    payload = dict(item)
    if isinstance(payload.get("standard_codes"), str):
        text = str(payload["standard_codes"]).strip()
        payload["standard_codes"] = [text] if text else []
    if isinstance(payload.get("sample_dimensions_mm"), (str, list)):
        dims = re.findall(r"-?\d+(?:\.\d+)?", str(payload["sample_dimensions_mm"]))
        numbers = [float(v) for v in dims]
        if numbers:
            payload.setdefault("sample_length_mm", numbers[0])
        if len(numbers) > 1:
            payload.setdefault("sample_width_mm", numbers[1])
        if len(numbers) > 2:
            payload.setdefault("sample_height_mm", numbers[2])
        payload.pop("sample_dimensions_mm", None)
    for field_name in ("sample_length_mm", "sample_width_mm", "sample_height_mm"):
        value = payload.get(field_name)
        if isinstance(value, str):
            dims = re.findall(r"-?\d+(?:\.\d+)?", value)
            payload[field_name] = float(dims[0]) if dims else None
    for field_name in (
        "pricing_quantity", "sample_count", "repeat_count", "sample_weight_kg",
        "required_temp_min", "required_temp_max",
        "required_humidity_min", "required_humidity_max",
        "required_temp_change_rate",
        "required_freq_min", "required_freq_max",
        "required_accel_min", "required_accel_max",
        "required_displacement_min", "required_displacement_max",
        "required_irradiance_min", "required_irradiance_max",
        "required_water_temp_min", "required_water_temp_max",
        "required_water_flow_min", "required_water_flow_max",
    ):
        value = payload.get(field_name)
        if isinstance(value, str):
            dims = re.findall(r"-?\d+(?:\.\d+)?", value)
            payload[field_name] = float(dims[0]) if dims else None
    _mirror_single_value_ranges(payload)
    return payload


def _mirror_single_value_ranges(payload: dict[str, Any]) -> None:
    for min_field, max_field in RANGE_FIELD_PAIRS:
        min_value = payload.get(min_field)
        max_value = payload.get(max_field)
        if min_value is None and max_value is None:
            continue
        if min_value is None and max_value is not None:
            payload[min_field] = max_value
        elif max_value is None and min_value is not None:
            payload[max_field] = min_value


class QwenRequester:
    def __init__(
        self,
        *,
        client: OpenAI | None = None,
        model: str | None = None,
        prompts_path: Path | None = None,
    ) -> None:
        settings = get_settings()
        self.model = model or settings.qwen_model
        self.prompts = _load_prompts(prompts_path or settings.prompts_path)
        self.client = client or OpenAI(api_key=settings.qwen_api_key, base_url=settings.qwen_base_url)

    def extract_form(
        self,
        documents: list[NormalizedDocument],
        *,
        test_type_options: list[str] | None = None,
        run_dir: Path | None = None,
    ) -> ModelFillResult:
        if not documents:
            return ModelFillResult(items=[], summary="未收到可抽取文档")
        prompt = self.prompts["document_extract"]
        messages = self._build_messages(
            system_prompt=str(prompt["system"]),
            user_template=str(prompt["user"]),
            documents=documents,
            current_rows=None,
            visible_fields=DOCUMENT_EXTRACT_VISIBLE_FIELDS,
            test_type_options=test_type_options or [],
        )
        content = self._stream_text(messages, run_dir=run_dir, request_name="文档抽取", max_tokens=12000)
        try:
            return self._parse_form_result(content)
        except json.JSONDecodeError:
            self._save_bad_json_response(content, run_dir=run_dir, prefix="document_extract")
            raise

    def enrich_form_from_documents(
        self,
        documents: list[NormalizedDocument],
        current_rows: list[FormRow],
        *,
        target_fields_by_row: dict[str, list[str]],
        run_dir: Path | None = None,
    ) -> ModelFillResult:
        if not documents:
            return ModelFillResult(items=[row.model_copy(deep=True) for row in current_rows], summary="未收到可补充文档")
        if not current_rows or not any(target_fields_by_row.values()):
            return ModelFillResult(items=[row.model_copy(deep=True) for row in current_rows], summary="无文档定向补充目标字段")
        prompt = self.prompts["document_targeted_enrich"]
        messages = self._build_document_enrich_messages(
            system_prompt=str(prompt["system"]),
            user_template=str(prompt["user"]),
            documents=documents,
            current_rows=current_rows,
            target_fields_by_row=target_fields_by_row,
            visible_fields=_document_targeted_visible_fields(target_fields_by_row),
        )
        content = self._stream_text(messages, run_dir=run_dir, request_name="文档定向补充", max_tokens=8000)
        try:
            return self._parse_form_result(content)
        except json.JSONDecodeError:
            self._save_bad_json_response(content, run_dir=run_dir, prefix="document_targeted_enrich")
            raise

    def split_batch_excel(self, document: NormalizedDocument, *, run_dir: Path | None = None) -> BatchQuoteSplitResult:
        prompt = self.prompts["batch_excel_split"]
        schema_json = json.dumps(
            {
                "items": [
                    {
                        "quote_id": "quote_1",
                        "title": "报价需求1",
                        "source_summary": "来源工作表/行号/识别依据",
                        "block_ids": ["Sheet1-row-2", "Sheet1-row-3"],
                    }
                ],
                "summary": "拆分摘要",
            },
            ensure_ascii=False,
            indent=2,
        )
        user_text = (
            str(prompt["user"])
            .replace("$document_manifest", self._document_manifest([document]))
            .replace("$document_text", self._batch_split_document_text(document))
            .replace("$schema_json", schema_json)
        )
        content = [{"type": "text", "text": user_text}]
        messages = [{"role": "system", "content": str(prompt["system"])}, {"role": "user", "content": content}]
        text = self._stream_text(messages, run_dir=run_dir, request_name="批量Excel拆分", max_tokens=8000)
        try:
            result = self._parse_batch_split_result(text)
        except json.JSONDecodeError:
            if run_dir is not None:
                (run_dir / "batch_split_raw_response.txt").write_text(text, encoding="utf-8")
            logger.exception("批量Excel拆分响应不是合法JSON，降级为整份Excel单子报价")
            result = BatchQuoteSplitResult(items=[], summary="批量拆分响应解析失败，按整份 Excel 作为一个报价需求处理", raw_response=text)
        if not result.items:
            result.items.append(
                BatchQuoteSplitItem(
                    quote_id="quote_1",
                    title=document.source_name or "报价需求1",
                    source_summary="模型未拆分出子报价，按整份 Excel 作为一个报价需求处理",
                    text="",
                    block_ids=[block.block_id for block in document.text_blocks],
                )
            )
        return result

    def split_batch_excel_with_protocol(self, workbook_path: Path, *, run_dir: Path | None = None) -> BatchExcelChunkSplitResult:
        from backend.quote.excel_protocol import ExcelChunkProtocol, ExcelReadProtocol

        prompt = self.prompts["batch_excel_protocol_split"]
        read_protocol = ExcelReadProtocol(workbook_path)
        chunk_protocol = ExcelChunkProtocol(workbook_path, sheet=read_protocol.get_current_sheet())
        workbook_summary = _excel_workbook_summary(workbook_path)
        user_text = str(prompt["user"]).replace("$workbook_summary", workbook_summary)
        messages: list[dict[str, Any]] = [
            {"role": "system", "content": str(prompt["system"])},
            {"role": "user", "content": [{"type": "text", "text": user_text}]},
        ]
        items: list[BatchExcelChunkItem] = []
        raw_responses: list[str] = []
        used_ids: set[str] = set()
        observed: dict[str, dict[str, set[int]]] = {}
        finished = False

        if run_dir is not None:
            append_run_log(run_dir, f"批量Excel协议切分开始: {workbook_path.name}")

        for turn in range(1, 13):
            content = self._stream_text(messages, run_dir=run_dir, request_name=f"批量Excel协议切分#{turn}", max_tokens=2500)
            raw_responses.append(content)
            if run_dir is not None:
                (run_dir / f"batch_excel_protocol_turn_{turn}.txt").write_text(content, encoding="utf-8")
            actions = _parse_protocol_actions(content)
            observations: list[dict[str, Any]] = []
            has_read_action = any(
                str(action_payload.get("action") or "").strip() in {"list_sheets", "set_sheet", "get_rows", "get_columns", "get_cell"}
                for action_payload in actions
            )
            created_before_turn = len(items)
            create_failed_this_turn = False

            if not actions:
                observations.append({"ok": False, "error": "No valid actions found. Output JSON with an actions array."})
            for action_payload in actions[:6]:
                action_name = str(action_payload.get("action") or "").strip()
                if has_read_action and action_name in {"create_chunk", "finish"}:
                    observations.append(
                        {
                            "ok": False,
                            "action": action_name,
                            "error": "Do not mix read actions with create_chunk/finish in the same turn. Review the observations first, then create chunks in the next turn.",
                        }
                    )
                    continue
                if action_name == "finish" and (len(items) < 2 or create_failed_this_turn):
                    observations.append(
                        {
                            "ok": False,
                            "action": action_name,
                            "error": (
                                "Cannot finish with fewer than 2 chunks. Batch quote mode always requires multiple "
                                "same-level quote chunks. Split the workbook into separate quote items; use shared_ranges "
                                "for titles, headers, product/group context, methods, criteria, and other shared rows."
                            ),
                        }
                    )
                    continue
                observation = self._execute_excel_protocol_action(
                    action_payload,
                    read_protocol=read_protocol,
                    chunk_protocol=chunk_protocol,
                    items=items,
                    used_ids=used_ids,
                    observed=observed,
                )
                observations.append(observation)
                if action_name == "create_chunk" and not observation.get("ok"):
                    create_failed_this_turn = True
                if observation.get("action") == "finish" and observation.get("ok"):
                    finished = True
                    break

            if create_failed_this_turn and len(items) > created_before_turn:
                observations.append(
                    {
                        "ok": True,
                        "action": "create_chunk",
                        "created_chunks": len(items),
                        "note": "Some create_chunk actions failed, but successful chunks in the same turn were kept. Retry only the failed ranges.",
                    }
                )

            if run_dir is not None:
                append_run_log(run_dir, f"批量Excel协议切分轮次 {turn}: actions={len(actions)} chunks={len(items)} finished={finished}")
            if finished:
                break

            messages.append({"role": "assistant", "content": content})
            messages.append({"role": "user", "content": [{"type": "text", "text": _protocol_observation_text(observations)}]})

        if len(items) < 2:
            if run_dir is not None:
                append_run_log(run_dir, f"批量Excel协议切分失败: chunks={len(items)}，批量模式禁止退化为单个报价")
            raise RuntimeError(
                "batch_excel_split_requires_multiple_chunks: 批量报价模式必须切分为多份同级报价，不能退化为单个报价或整表报价。"
            )

        summary = f"协议切分生成 {len(items)} 个子报价" if finished else f"协议切分达到轮次上限，已生成 {len(items)} 个子报价"
        if run_dir is not None:
            append_run_log(run_dir, f"批量Excel协议切分完成: chunks={len(items)} finished={finished}")
        return BatchExcelChunkSplitResult(items=items, summary=summary, raw_response="\n\n".join(raw_responses))

    def _execute_excel_protocol_action(
        self,
        payload: dict[str, Any],
        *,
        read_protocol: Any,
        chunk_protocol: Any,
        items: list[BatchExcelChunkItem],
        used_ids: set[str],
        observed: dict[str, dict[str, set[int]]],
    ) -> dict[str, Any]:
        action = str(payload.get("action") or "").strip()
        try:
            if action == "list_sheets":
                return {"ok": True, "action": action, "sheets": read_protocol.list_sheets(), "current_sheet": read_protocol.get_current_sheet()}
            if action == "set_sheet":
                sheet = str(payload.get("sheet") or "").strip()
                read_protocol.set_current_sheet(sheet)
                chunk_protocol.set_current_sheet(sheet)
                return {"ok": True, "action": action, "current_sheet": sheet}
            if action == "get_rows":
                cells = read_protocol.get_rows(int(payload.get("start")), int(payload.get("end")))
                _mark_observed_rows(observed, read_protocol.get_current_sheet(), int(payload.get("start")), int(payload.get("end")))
                return {"ok": True, "action": action, "current_sheet": read_protocol.get_current_sheet(), "cells": _serialize_cells(cells)}
            if action == "get_columns":
                cells = read_protocol.get_columns(str(payload.get("start") or ""), str(payload.get("end") or ""))
                _mark_observed_columns(observed, read_protocol.get_current_sheet(), str(payload.get("start") or ""), str(payload.get("end") or ""))
                return {"ok": True, "action": action, "current_sheet": read_protocol.get_current_sheet(), "cells": _serialize_cells(cells)}
            if action == "get_cell":
                cell = read_protocol.get_cell(str(payload.get("address") or ""))
                _mark_observed_cell(observed, read_protocol.get_current_sheet(), str(payload.get("address") or ""))
                return {"ok": True, "action": action, "current_sheet": read_protocol.get_current_sheet(), "cell": cell.to_compact_dict()}
            if action == "create_chunk":
                shared_ranges = _normalize_string_list(payload.get("shared_ranges"))
                main_range = _shrink_main_trailing_shared_rows(str(payload.get("main_range") or ""), shared_ranges)
                current_sheet = chunk_protocol.get_current_sheet()
                missing_ref = _first_unobserved_range(
                    observed,
                    current_sheet,
                    [
                        main_range,
                        *shared_ranges,
                    ],
                )
                if missing_ref:
                    return {
                        "ok": False,
                        "action": action,
                        "error": f'Range "{missing_ref}" has not been sufficiently observed. Use get_rows for the row span or get_columns for the column span before create_chunk.',
                    }
                oversized_error = _oversized_batch_chunk_error(main_range, shared_ranges, chunk_protocol)
                if oversized_error:
                    return {"ok": False, "action": action, "error": oversized_error}
                chunk = chunk_protocol.create_chunk(
                    main_range=main_range,
                    shared_ranges=shared_ranges,
                    label=str(payload.get("label") or "").strip() or None,
                )
                overlap_ratio = _shared_main_overlap_ratio(chunk.main_range, chunk.shared_ranges)
                if overlap_ratio > 0.5:
                    return {
                        "ok": False,
                        "action": action,
                        "error": (
                            "shared_ranges overlap too much with main_range. "
                            "Keep product/sample/project-specific rows or columns in main_range, and put only truly shared title/notes/method rows or columns in shared_ranges."
                        ),
                    }
                chunk_key = (chunk_protocol.get_current_sheet(), chunk.main_range, tuple(chunk.shared_ranges))
                existing_keys = {
                    (
                        str(existing.source_summary).split(" main_range=", 1)[0].replace("sheet=", ""),
                        existing.chunk.main_range,
                        tuple(existing.chunk.shared_ranges),
                    )
                    for existing in items
                }
                if chunk_key in existing_keys:
                    return {
                        "ok": False,
                        "action": action,
                        "error": f"Duplicate chunk range ignored: main_range={chunk.main_range}, shared_ranges={chunk.shared_ranges}",
                    }
                index = len(items) + 1
                title = chunk.label or f"报价需求{index}"
                quote_id = _normalize_quote_id(str(payload.get("quote_id") or title), index, used_ids)
                used_ids.add(quote_id)
                items.append(
                    BatchExcelChunkItem(
                        quote_id=quote_id,
                        title=title,
                        source_summary=f"sheet={chunk_protocol.get_current_sheet()} main_range={chunk.main_range} shared_ranges={','.join(chunk.shared_ranges) or '-'}",
                        chunk=chunk,
                    )
                )
                return {
                    "ok": True,
                    "action": action,
                    "quote_id": quote_id,
                    "label": title,
                    "main_range": chunk.main_range,
                    "shared_ranges": chunk.shared_ranges,
                    "created_chunks": len(items),
                }
            if action == "finish":
                return {"ok": True, "action": action, "summary": str(payload.get("summary") or "").strip(), "created_chunks": len(items)}
            return {"ok": False, "action": action or "(missing)", "error": "Unknown action."}
        except Exception as exc:
            return {"ok": False, "action": action or "(missing)", "error": str(exc)}

    def document_from_batch_item(self, source: NormalizedDocument, item: BatchQuoteSplitItem) -> NormalizedDocument:
        blocks_by_id = {block.block_id: block for block in source.text_blocks}
        selected_blocks = [blocks_by_id[block_id] for block_id in item.block_ids if block_id in blocks_by_id]
        if selected_blocks:
            text_blocks = [
                NormalizedTextBlock(
                    block_id=f"{item.quote_id}-{block.block_id}",
                    block_type=block.block_type,
                    text=block.text,
                    source_path=block.source_path,
                )
                for block in selected_blocks
            ]
        else:
            text_blocks = [
                NormalizedTextBlock(
                    block_id=f"{item.quote_id}-batch-split",
                    block_type="BatchQuote",
                    text=item.text or self._document_text([source]),
                    source_path=item.source_summary,
                )
            ]
        return NormalizedDocument(
            document_id=item.quote_id,
            source_name=item.title or item.quote_id,
            source_kind="batch_excel_quote",
            original_path=source.original_path,
            text_blocks=text_blocks,
            assets=source.assets,
            metadata={**source.metadata, "batch_quote_id": item.quote_id, "batch_source_summary": item.source_summary},
        )

    def enrich_form_with_evidences(
        self,
        current_rows: list[FormRow],
        *,
        target_fields_by_row: dict[str, list[str]] | None = None,
        run_dir: Path | None = None,
    ) -> ModelFillResult:
        if not any(row.standard_evidences for row in current_rows):
            return ModelFillResult(items=[row.model_copy(deep=True) for row in current_rows], summary="无可补充标准证据")
        prompt = self.prompts["standard_enrich_with_evidences"]
        messages = self._build_evidence_messages(
            system_prompt=str(prompt["system"]),
            user_template=str(prompt["user"]),
            current_rows=current_rows,
            target_fields_by_row=target_fields_by_row or {},
            visible_fields=STANDARD_ENRICH_VISIBLE_FIELDS,
        )
        content = self._stream_text(messages, run_dir=run_dir, request_name="标准证据补表", max_tokens=8000)
        try:
            return self._parse_form_result(content)
        except json.JSONDecodeError:
            self._save_bad_json_response(content, run_dir=run_dir, prefix="standard_enrich")
            raise

    def discover_standard_fields(
        self,
        current_rows: list[FormRow],
        *,
        target_fields_by_row: dict[str, list[str]] | None = None,
        supported_fields: list[str] | None = None,
        run_dir: Path | None = None,
    ) -> StandardFieldDiscoveryResult:
        if not any(row.standard_evidences for row in current_rows):
            return StandardFieldDiscoveryResult(items=[], summary="无可用于字段发现的标准证据")
        prompt = self.prompts["standard_field_discovery"]
        messages = self._build_discovery_messages(
            system_prompt=str(prompt["system"]),
            user_template=str(prompt["user"]),
            current_rows=current_rows,
            target_fields_by_row=target_fields_by_row or {},
            supported_fields=supported_fields or [],
            visible_fields=STANDARD_DISCOVERY_VISIBLE_FIELDS,
        )
        content = self._stream_text(messages, run_dir=run_dir, request_name="标准字段发现")
        return self._parse_standard_field_discovery_result(content, supported_fields=supported_fields or [])

    def judge_standard_context(
        self,
        row: FormRow,
        evidence: StandardEvidence,
        *,
        target_fields: list[str] | None = None,
        run_dir: Path | None = None,
    ) -> StandardContextDecision:
        prompt = self.prompts["standard_context_judge"]
        messages = self._build_context_judge_messages(
            system_prompt=str(prompt["system"]),
            user_template=str(prompt["user"]),
            row=row,
            evidence=evidence,
            target_fields=target_fields or [],
            visible_fields=CONTEXT_JUDGE_VISIBLE_FIELDS,
        )
        content = self._stream_text(messages, run_dir=run_dir, request_name="标准上下文判定")
        return self._parse_standard_context_decision(content)

    def _build_messages(
        self,
        *,
        system_prompt: str,
        user_template: str,
        documents: list[NormalizedDocument],
        current_rows: list[FormRow] | None,
        visible_fields: tuple[str, ...],
        test_type_options: list[str] | None = None,
    ) -> list[dict[str, Any]]:
        manifest = self._document_manifest(documents)
        document_context = self._document_context(documents)
        document_text = self._document_text(documents)
        test_type_options_text = _test_type_options_text(test_type_options or [])
        schema_json = json.dumps(_schema_example(visible_fields), ensure_ascii=False, indent=2)
        current_form = json.dumps(self._rows_for_model(current_rows or [], visible_fields), ensure_ascii=False, indent=2)
        user_text = (
            user_template.replace("$document_manifest", manifest)
            .replace("$document_text", document_text)
            .replace("$test_type_options", test_type_options_text)
            .replace("$schema_json", schema_json)
            .replace(
                "$rules_text",
                _rules_text(
                    preserve_row_ids=current_rows is not None,
                    include_sample_count=True,
                ),
            )
            .replace("$current_form", current_form)
        )
        if document_context:
            user_text = self._insert_document_context(user_text, document_context)

        content: list[dict[str, Any]] = []
        for document in documents:
            for asset in document.assets:
                content.append({"type": "image_url", "image_url": {"url": asset.data_url}})
        content.append({"type": "text", "text": user_text})
        return [{"role": "system", "content": system_prompt}, {"role": "user", "content": content}]

    def _build_document_enrich_messages(
        self,
        *,
        system_prompt: str,
        user_template: str,
        documents: list[NormalizedDocument],
        current_rows: list[FormRow],
        target_fields_by_row: dict[str, list[str]],
        visible_fields: tuple[str, ...],
    ) -> list[dict[str, Any]]:
        manifest = self._document_manifest(documents)
        document_context = self._document_context(documents)
        document_text = self._document_text(documents)
        current_form = json.dumps(self._rows_for_model(current_rows, visible_fields), ensure_ascii=False, indent=2)
        target_manifest, target_text = self._row_target_fields_text(
            current_rows,
            target_fields_by_row,
            empty_text="(无文档定向补充目标字段)",
            intro_text="仅允许补充或修正以下字段：",
        )
        schema_json = json.dumps(_schema_example(visible_fields), ensure_ascii=False, indent=2)
        user_text = (
            user_template.replace("$document_manifest", manifest)
            .replace("$document_text", document_text)
            .replace("$current_form", current_form)
            .replace("$target_manifest", target_manifest)
            .replace("$target_text", target_text)
            .replace("$schema_json", schema_json)
            .replace(
                "$rules_text",
                _rules_text(
                    preserve_row_ids=True,
                    include_sample_count=True,
                ),
            )
        )
        if document_context:
            user_text = self._insert_document_context(user_text, document_context)

        content: list[dict[str, Any]] = []
        for document in documents:
            for asset in document.assets:
                content.append({"type": "image_url", "image_url": {"url": asset.data_url}})
        content.append({"type": "text", "text": user_text})
        return [{"role": "system", "content": system_prompt}, {"role": "user", "content": content}]

    def _build_evidence_messages(
        self,
        *,
        system_prompt: str,
        user_template: str,
        current_rows: list[FormRow],
        target_fields_by_row: dict[str, list[str]],
        visible_fields: tuple[str, ...],
    ) -> list[dict[str, Any]]:
        schema_json = json.dumps(_schema_example(visible_fields), ensure_ascii=False, indent=2)
        current_form = json.dumps(self._rows_for_model(current_rows, visible_fields), ensure_ascii=False, indent=2)
        evidence_manifest, evidence_text = self._row_evidence_text(current_rows)
        target_manifest, target_text = self._row_target_fields_text(current_rows, target_fields_by_row)
        user_text = (
            user_template.replace("$current_form", current_form)
            .replace("$target_manifest", target_manifest)
            .replace("$target_text", target_text)
            .replace("$evidence_manifest", evidence_manifest)
            .replace("$evidence_text", evidence_text)
            .replace("$schema_json", schema_json)
            .replace(
                "$rules_text",
                _rules_text(
                    preserve_row_ids=True,
                    include_sample_count=False,
                ),
            )
        )
        return [{"role": "system", "content": system_prompt}, {"role": "user", "content": [{"type": "text", "text": user_text}]}]

    def _build_context_judge_messages(
        self,
        *,
        system_prompt: str,
        user_template: str,
        row: FormRow,
        evidence: StandardEvidence,
        target_fields: list[str],
        visible_fields: tuple[str, ...],
    ) -> list[dict[str, Any]]:
        current_row = json.dumps(self._rows_for_model([row], visible_fields), ensure_ascii=False, indent=2)
        evidence_text = "\n".join(self._format_evidence_block(evidence))
        target_text = "\n".join(f"- {field}" for field in target_fields) if target_fields else "- (无明确目标字段)"
        user_text = (
            user_template.replace("$current_row", current_row)
            .replace("$evidence_text", evidence_text)
            .replace("$target_fields", target_text)
        )
        return [{"role": "system", "content": system_prompt}, {"role": "user", "content": [{"type": "text", "text": user_text}]}]

    def _build_discovery_messages(
        self,
        *,
        system_prompt: str,
        user_template: str,
        current_rows: list[FormRow],
        target_fields_by_row: dict[str, list[str]],
        supported_fields: list[str],
        visible_fields: tuple[str, ...],
    ) -> list[dict[str, Any]]:
        current_form = json.dumps(self._rows_for_model(current_rows, visible_fields), ensure_ascii=False, indent=2)
        evidence_manifest, evidence_text = self._row_evidence_text(current_rows)
        target_manifest, target_text = self._row_target_fields_text(current_rows, target_fields_by_row)
        supported_fields_text = "\n".join(f"- {field}" for field in supported_fields) if supported_fields else "- (无)"
        schema_json = json.dumps(
            {
                "items": [
                    {
                        "row_id": "",
                        "discovered_standard_fields": ["required_temp_max"],
                        "extra_standard_requirements": [
                            {
                                "requirement_name": "通电状态",
                                "requirement_text": "试验期间样品应保持通电运行",
                                "source_section": "5.1.3",
                            }
                        ],
                    }
                ]
            },
            ensure_ascii=False,
            indent=2,
        )
        user_text = (
            user_template.replace("$current_form", current_form)
            .replace("$target_manifest", target_manifest)
            .replace("$target_text", target_text)
            .replace("$supported_fields_text", supported_fields_text)
            .replace("$evidence_manifest", evidence_manifest)
            .replace("$evidence_text", evidence_text)
            .replace("$schema_json", schema_json)
        )
        return [{"role": "system", "content": system_prompt}, {"role": "user", "content": [{"type": "text", "text": user_text}]}]

    def _document_manifest(self, documents: list[NormalizedDocument]) -> str:
        lines = [
            f"- {doc.source_name} | kind={doc.source_kind} | blocks={len(doc.text_blocks)} | images={len(doc.assets)}"
            for doc in documents
        ]
        return "\n".join(lines) if lines else "- 无"

    def _document_context(self, documents: list[NormalizedDocument]) -> str:
        sections: list[str] = []
        for document in documents:
            metadata = document.metadata or {}
            hints = _normalize_string_list(metadata.get("extraction_hints"))
            if not hints:
                continue

            lines = [f"## {document.source_name}"]
            chunk_label = str(metadata.get("chunk_label") or "").strip()
            main_range = str(metadata.get("main_range") or "").strip()
            shared_ranges = _normalize_string_list(metadata.get("shared_ranges"))
            if chunk_label:
                lines.append(f"- chunk_label: {chunk_label}")
            if main_range:
                lines.append(f"- main_range: {main_range}")
            if shared_ranges:
                lines.append(f"- shared_ranges: {', '.join(shared_ranges)}")
            lines.extend(f"- {hint}" for hint in hints)
            sections.append("\n".join(lines).strip())
        return "\n\n".join(sections).strip()

    def _insert_document_context(self, user_text: str, document_context: str) -> str:
        context_block = f"\n\n文档处理说明：\n{document_context}"
        marker = "\n\n结构化正文："
        if marker in user_text:
            return user_text.replace(marker, f"{context_block}{marker}", 1)
        return f"{context_block.lstrip()}\n\n{user_text}"

    def _document_text(self, documents: list[NormalizedDocument]) -> str:
        sections: list[str] = []
        for document in documents:
            block_text = "\n".join(block.text for block in document.text_blocks if block.text.strip()).strip()
            asset_lines = [
                f"- [{asset.asset_id}] position={asset.position} context={asset.context_text or '无'}"
                for asset in document.assets
            ]
            section = [f"## {document.source_name}", "正文：", block_text or "(空)"]
            if asset_lines:
                section.extend(["插图说明：", *asset_lines])
            sections.append("\n".join(section).strip())
        return "\n\n".join(sections).strip()

    def _batch_split_document_text(self, document: NormalizedDocument) -> str:
        sections: list[str] = []
        for block in document.text_blocks:
            text = block.text.strip()
            if not text:
                continue
            sections.append(f"### block_id={block.block_id} | source={block.source_path or '-'}\n{text}")
        return "\n\n".join(sections).strip()

    def _rows_for_model(self, rows: list[FormRow], visible_fields: tuple[str, ...]) -> dict[str, Any]:
        return {"items": [{field: _jsonable_value(getattr(row, field)) for field in visible_fields} for row in rows]}

    def _row_evidence_text(self, rows: list[FormRow]) -> tuple[str, str]:
        manifest_lines: list[str] = []
        sections: list[str] = []
        for row in rows:
            evidences = row.standard_evidences
            row_label = row.raw_test_type or row.canonical_test_type or row.row_id
            manifest_lines.append(f"- row_id={row.row_id} | test={row_label} | evidence_count={len(evidences)}")
            section_lines = [f"## row_id={row.row_id} | test={row_label}"]
            if not evidences:
                section_lines.append("(无标准证据)")
                sections.append("\n".join(section_lines))
                continue
            for evidence in evidences:
                section_lines.extend(self._format_evidence_block(evidence))
            sections.append("\n".join(section_lines).strip())
        manifest = "\n".join(manifest_lines) if manifest_lines else "- 无"
        return manifest, "\n\n".join(sections).strip()

    def _row_target_fields_text(
        self,
        rows: list[FormRow],
        target_fields_by_row: dict[str, list[str]],
        *,
        empty_text: str = "(无标准补充目标字段)",
        intro_text: str = "仅允许补充以下字段：",
    ) -> tuple[str, str]:
        manifest_lines: list[str] = []
        sections: list[str] = []
        for row in rows:
            targets = target_fields_by_row.get(row.row_id) or []
            manifest_lines.append(f"- row_id={row.row_id} | target_fields={','.join(targets) if targets else '(无)'}")
            label = row.raw_test_type or row.canonical_test_type or row.row_id
            lines = [f"## row_id={row.row_id} | test={label}"]
            if not targets:
                lines.append(empty_text)
            else:
                lines.extend([intro_text, *[f"- {field}" for field in targets]])
            sections.append("\n".join(lines))
        manifest = "\n".join(manifest_lines) if manifest_lines else "- 无"
        return manifest, "\n\n".join(sections).strip()

    def _format_evidence_block(self, evidence: StandardEvidence) -> list[str]:
        reasons = "、".join(evidence.match_reasons) if evidence.match_reasons else "无"
        return [
            f"### {evidence.standard_code} | {evidence.section_id or evidence.chunk_id} | pages={evidence.page_start}-{evidence.page_end}",
            f"标题：{evidence.section_title or '无'}",
            f"匹配原因：{reasons}",
            "证据正文：",
            evidence.text.strip() or "(空)",
        ]

    def _stream_text(self, messages: list[dict[str, Any]], *, run_dir: Path | None, request_name: str, max_tokens: int = 4000) -> str:
        user_content = messages[1]["content"]
        image_count = sum(1 for item in user_content if item.get("type") == "image_url")
        text_length = len(str(user_content[-1].get("text") or "")) if user_content else 0
        logger.info("开始模型请求: stage=%s model=%s images=%s text_chars=%s", request_name, self.model, image_count, text_length)
        if run_dir is not None:
            append_run_log(run_dir, f"开始模型请求: {request_name} | model={self.model} | images={image_count} | text_chars={text_length}")
            prompt_path = self._save_model_prompt(
                messages,
                run_dir=run_dir,
                request_name=request_name,
                max_tokens=max_tokens,
                image_count=image_count,
                text_length=text_length,
            )
            append_run_log(run_dir, f"模型Prompt已保存: {prompt_path.relative_to(run_dir)}")
        completion = self.client.chat.completions.create(
            model=self.model,
            messages=messages,
            temperature=0.1,
            max_tokens=max_tokens,
            modalities=["text"],
            stream=True,
            stream_options={"include_usage": True},
        )
        parts: list[str] = []
        chunk_count = 0
        for chunk in completion:
            if not getattr(chunk, "choices", None):
                continue
            delta = chunk.choices[0].delta
            content = getattr(delta, "content", None)
            if isinstance(content, str):
                parts.append(content)
                chunk_count += 1
            elif isinstance(content, list):
                parts.append(_flatten_stream_part(content))
                chunk_count += 1
        text = "".join(part for part in parts if part)
        if not text.strip():
            logger.error("模型请求为空响应: stage=%s model=%s", request_name, self.model)
            if run_dir is not None:
                append_run_log(run_dir, f"模型请求失败: {request_name} | empty_response")
            raise RuntimeError("qwen_empty_response")
        logger.info("模型请求成功: stage=%s chunks=%s response_chars=%s", request_name, chunk_count, len(text))
        if run_dir is not None:
            append_run_log(run_dir, f"模型请求成功: {request_name} | chunks={chunk_count} | response_chars={len(text)}")
        return text

    def _save_model_prompt(
        self,
        messages: list[dict[str, Any]],
        *,
        run_dir: Path,
        request_name: str,
        max_tokens: int,
        image_count: int,
        text_length: int,
    ) -> Path:
        prompts_dir = run_dir / "model_prompts"
        prompts_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d%H%M%S%f")
        safe_name = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff._-]+", "_", request_name).strip("_") or "model_request"
        path = prompts_dir / f"{timestamp}_{safe_name}.txt"
        path.write_text(
            _render_model_prompt_log(
                messages,
                request_name=request_name,
                model=self.model,
                max_tokens=max_tokens,
                image_count=image_count,
                text_length=text_length,
            ),
            encoding="utf-8",
        )
        return path

    def _save_bad_json_response(self, content: str, *, run_dir: Path | None, prefix: str) -> None:
        if run_dir is None:
            return
        filename = f"{prefix}_bad_json_{datetime.now().strftime('%Y%m%d%H%M%S')}.txt"
        (run_dir / filename).write_text(content, encoding="utf-8")
        append_run_log(run_dir, f"模型JSON解析失败，原始响应已保存: {filename}")

    def _parse_form_result(self, content: str) -> ModelFillResult:
        payload = json.loads(_extract_json_text(content))
        if isinstance(payload, list):
            raw_items, summary = payload, ""
        else:
            raw_items = payload.get("items") or []
            summary = str(payload.get("document_summary") or payload.get("summary") or "").strip()

        items: list[FormRow] = []
        for item in raw_items:
            if not isinstance(item, dict):
                continue
            items.append(FormRow.model_validate(_normalize_item_payload(item)))
        logger.info("模型响应解析完成: rows=%s summary=%s", len(items), summary or "-")
        return ModelFillResult(items=items, summary=summary, raw_response=content)

    def _parse_batch_split_result(self, content: str) -> BatchQuoteSplitResult:
        payload = json.loads(_extract_json_text(content))
        if isinstance(payload, list):
            raw_items, summary = payload, ""
        elif isinstance(payload, dict):
            raw_items = payload.get("items") or []
            summary = str(payload.get("summary") or "").strip()
        else:
            raw_items, summary = [], ""

        items: list[BatchQuoteSplitItem] = []
        used_ids: set[str] = set()
        for index, raw in enumerate(raw_items, start=1):
            if not isinstance(raw, dict):
                continue
            quote_id = _normalize_quote_id(str(raw.get("quote_id") or ""), index, used_ids)
            title = str(raw.get("title") or f"报价需求{index}").strip() or f"报价需求{index}"
            source_summary = str(raw.get("source_summary") or "").strip()
            text = str(raw.get("text") or "").strip()
            block_ids = _normalize_string_list(raw.get("block_ids") or raw.get("source_block_ids"))
            if not text and not block_ids:
                continue
            used_ids.add(quote_id)
            items.append(BatchQuoteSplitItem(quote_id=quote_id, title=title, source_summary=source_summary, text=text, block_ids=block_ids))
        logger.info("批量Excel拆分解析完成: quotes=%s summary=%s", len(items), summary or "-")
        return BatchQuoteSplitResult(items=items, summary=summary, raw_response=content)

    def _parse_standard_context_decision(self, content: str) -> StandardContextDecision:
        payload = json.loads(_extract_json_text(content))
        if not isinstance(payload, dict):
            raise RuntimeError("invalid_standard_context_decision")
        decision = StandardContextDecision.model_validate(payload)
        logger.info("标准上下文判定完成: decision=%s reason=%s", decision.decision, decision.reason or "-")
        return decision

    def _parse_standard_field_discovery_result(
        self,
        content: str,
        *,
        supported_fields: list[str],
    ) -> StandardFieldDiscoveryResult:
        payload = json.loads(_extract_json_text(content))
        if isinstance(payload, list):
            raw_items, summary = payload, ""
        else:
            raw_items = payload.get("items") or []
            summary = str(payload.get("summary") or "").strip()

        supported = set(supported_fields)
        items: list[StandardFieldDiscoveryItem] = []
        for item in raw_items:
            if not isinstance(item, dict):
                continue
            row_id = str(item.get("row_id") or "").strip()
            discovered = _normalize_discovered_fields(item.get("discovered_standard_fields"), supported)
            extras = _normalize_extra_requirements(item.get("extra_standard_requirements"))
            if row_id:
                items.append(
                    StandardFieldDiscoveryItem(
                        row_id=row_id,
                        discovered_standard_fields=discovered,
                        extra_standard_requirements=extras,
                    )
                )
        logger.info("标准字段发现解析完成: rows=%s summary=%s", len(items), summary or "-")
        return StandardFieldDiscoveryResult(items=items, summary=summary, raw_response=content)


def _excel_workbook_summary(workbook_path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    lines = [f"- file={workbook_path.name}", f"- sheet_count={len(workbook.sheetnames)}"]
    for sheet in workbook.worksheets:
        lines.append(
            f"- sheet={sheet.title} | max_row={sheet.max_row} | max_col={sheet.max_column} "
            f"(A-{get_column_letter(sheet.max_column)})"
        )
    workbook.close()
    return "\n".join(lines)


def _render_model_prompt_log(
    messages: list[dict[str, Any]],
    *,
    request_name: str,
    model: str,
    max_tokens: int,
    image_count: int,
    text_length: int,
) -> str:
    lines = [
        f"request_name: {request_name}",
        f"model: {model}",
        "temperature: 0.1",
        f"max_tokens: {max_tokens}",
        "modalities: text",
        f"image_count: {image_count}",
        f"text_chars: {text_length}",
        "",
    ]
    for message_index, message in enumerate(messages, start=1):
        role = str(message.get("role") or "")
        lines.append(f"===== message {message_index} | role={role} =====")
        content = message.get("content")
        if isinstance(content, str):
            lines.append(content)
            lines.append("")
            continue
        if isinstance(content, list):
            for part_index, part in enumerate(content, start=1):
                if not isinstance(part, dict):
                    lines.append(f"--- part {part_index} | raw ---")
                    lines.append(str(part))
                    continue
                part_type = str(part.get("type") or "")
                lines.append(f"--- part {part_index} | type={part_type} ---")
                if part_type == "text":
                    lines.append(str(part.get("text") or ""))
                    continue
                if part_type == "image_url":
                    image_url = part.get("image_url")
                    url = ""
                    if isinstance(image_url, dict):
                        url = str(image_url.get("url") or "")
                    lines.append(f"[image_url omitted] chars={len(url)} prefix={url[:48]}")
                    continue
                lines.append(json.dumps(part, ensure_ascii=False, default=str))
            lines.append("")
            continue
        lines.append(str(content))
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def _parse_protocol_actions(content: str) -> list[dict[str, Any]]:
    json_text = _extract_json_text(content)
    try:
        payload = json.loads(json_text)
    except json.JSONDecodeError:
        try:
            payload = json.loads(_repair_protocol_json(json_text))
        except json.JSONDecodeError:
            return []
    raw_actions: Any
    if isinstance(payload, dict) and isinstance(payload.get("actions"), list):
        raw_actions = payload["actions"]
    elif isinstance(payload, dict) and payload.get("action"):
        raw_actions = [payload]
    elif isinstance(payload, list):
        raw_actions = payload
    else:
        raw_actions = []
    return [item for item in raw_actions if isinstance(item, dict)]


def _repair_protocol_json(text: str) -> str:
    repaired = text.strip()
    repaired = re.sub(r'("end"\s*:\s*"[^"]+")\s*\]', r"\1}", repaired)
    repaired = re.sub(r'("start"\s*:\s*"[^"]+")\s*\]', r"\1}", repaired)
    return repaired


def _observed_bucket(observed: dict[str, dict[str, set[int]]], sheet: str) -> dict[str, set[int]]:
    return observed.setdefault(sheet, {"rows": set(), "cols": set()})


def _mark_observed_rows(observed: dict[str, dict[str, set[int]]], sheet: str, start: int, end: int) -> None:
    if start > end:
        return
    _observed_bucket(observed, sheet)["rows"].update(range(max(1, start), end + 1))


def _mark_observed_columns(observed: dict[str, dict[str, set[int]]], sheet: str, start: str, end: str) -> None:
    try:
        start_col = column_index_from_string(str(start or "").strip().upper())
        end_col = column_index_from_string(str(end or "").strip().upper())
    except ValueError:
        return
    if start_col > end_col:
        return
    _observed_bucket(observed, sheet)["cols"].update(range(start_col, end_col + 1))


def _mark_observed_cell(observed: dict[str, dict[str, set[int]]], sheet: str, address: str) -> None:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(address)
    except ValueError:
        return
    _observed_bucket(observed, sheet)["rows"].update(range(min_row, max_row + 1))
    _observed_bucket(observed, sheet)["cols"].update(range(min_col, max_col + 1))


def _first_unobserved_range(observed: dict[str, dict[str, set[int]]], sheet: str, refs: list[str]) -> str:
    bucket = observed.get(sheet) or {"rows": set(), "cols": set()}
    observed_rows = bucket["rows"]
    observed_cols = bucket["cols"]
    for ref in refs:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(ref)
        except ValueError:
            continue
        row_span_observed = all(row in observed_rows for row in range(min_row, max_row + 1))
        col_span_observed = all(col in observed_cols for col in range(min_col, max_col + 1))
        if not row_span_observed and not col_span_observed:
            return ref
    return ""


def _shared_main_overlap_ratio(main_range: str, shared_ranges: list[str]) -> float:
    try:
        main_min_col, main_min_row, main_max_col, main_max_row = range_boundaries(main_range)
    except ValueError:
        return 0.0
    main_area = (main_max_col - main_min_col + 1) * (main_max_row - main_min_row + 1)
    if main_area <= 0:
        return 0.0
    overlap_cells: set[tuple[int, int]] = set()
    for ref in shared_ranges:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(ref)
        except ValueError:
            continue
        for row in range(max(main_min_row, min_row), min(main_max_row, max_row) + 1):
            for col in range(max(main_min_col, min_col), min(main_max_col, max_col) + 1):
                overlap_cells.add((row, col))
    return len(overlap_cells) / main_area


def _oversized_batch_chunk_error(main_range: str, shared_ranges: list[str], chunk_protocol: Any) -> str:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(main_range)
        sheet = chunk_protocol._sheet()
    except Exception:
        return ""
    sheet_area = max(1, int(sheet.max_row or 1) * int(sheet.max_column or 1))
    main_area = (max_col - min_col + 1) * (max_row - min_row + 1)
    row_ratio = (max_row - min_row + 1) / max(1, int(sheet.max_row or 1))
    col_ratio = (max_col - min_col + 1) / max(1, int(sheet.max_column or 1))
    covers_most_sheet = main_area / sheet_area >= 0.55
    spans_many_rows_and_cols = (max_row - min_row + 1) >= max(4, int((sheet.max_row or 1) * 0.5)) and (
        max_col - min_col + 1
    ) >= max(4, int((sheet.max_column or 1) * 0.5))
    spans_full_width_block = col_ratio >= 0.75 and row_ratio >= 0.5 and (max_row - min_row + 1) >= 4
    if not (covers_most_sheet and spans_many_rows_and_cols) and not spans_full_width_block:
        return ""
    shared_text = ", ".join(shared_ranges) if shared_ranges else "(none)"
    return (
        f'main_range "{main_range}" covers most of the sheet and is likely merging multiple independent quote items. '
        "Batch quote mode must split the workbook into multiple same-level quote chunks; do not create a single full-table/full-series chunk. "
        "Keep each quote-specific product area or test row in its own main_range, and move shared titles, headers, methods, criteria, "
        f"addresses, group/product context, and other common rows into shared_ranges. Current shared_ranges={shared_text}."
    )


def _shrink_main_trailing_shared_rows(main_range: str, shared_ranges: list[str]) -> str:
    try:
        main_min_col, main_min_row, main_max_col, main_max_row = range_boundaries(main_range)
    except ValueError:
        return main_range
    adjusted_max_row = main_max_row
    for ref in shared_ranges:
        try:
            shared_min_col, shared_min_row, shared_max_col, shared_max_row = range_boundaries(ref)
        except ValueError:
            continue
        covers_main_columns = shared_min_col <= main_min_col and main_max_col <= shared_max_col
        overlaps_trailing_rows = main_min_row < shared_min_row <= adjusted_max_row <= shared_max_row
        if covers_main_columns and overlaps_trailing_rows:
            adjusted_max_row = shared_min_row - 1
    if adjusted_max_row == main_max_row or adjusted_max_row < main_min_row:
        return main_range
    return f"{get_column_letter(main_min_col)}{main_min_row}:{get_column_letter(main_max_col)}{adjusted_max_row}"


def _protocol_observation_text(observations: list[dict[str, Any]]) -> str:
    payload = {"observations": observations}
    return (
        "工具执行结果如下。请继续观察或创建 chunk；如果已创建完所有独立报价需求，请输出 finish。\n"
        f"{json.dumps(payload, ensure_ascii=False, default=str)}"
    )


def _serialize_cells(cells: list[Any], *, limit: int = 500) -> dict[str, Any]:
    serialized: list[dict[str, Any]] = []
    omitted_empty_cells = 0
    truncated_cells = 0
    for cell in cells:
        if getattr(cell, "value", None) is None and not getattr(cell, "is_merged", False):
            omitted_empty_cells += 1
            continue
        if len(serialized) >= limit:
            truncated_cells += 1
            continue
        serialized.append(cell.to_compact_dict())
    return {
        "items": serialized,
        "returned": len(serialized),
        "omitted_empty_cells": omitted_empty_cells,
        "truncated_cells": truncated_cells,
    }


def _normalize_discovered_fields(value: Any, supported_fields: set[str]) -> list[str]:
    raw_values: list[str]
    if isinstance(value, list):
        raw_values = [str(item).strip() for item in value]
    elif isinstance(value, str):
        raw_values = [part.strip() for part in re.split(r"[,，;\n]+", value) if part.strip()]
    else:
        raw_values = []

    result: list[str] = []
    seen: set[str] = set()
    for field_name in raw_values:
        if field_name in supported_fields and field_name not in seen:
            seen.add(field_name)
            result.append(field_name)
    return result


def _normalize_extra_requirements(value: Any) -> list[ExtraStandardRequirement]:
    if not isinstance(value, list):
        return []
    items: list[ExtraStandardRequirement] = []
    seen: set[tuple[str, str, str]] = set()
    for raw in value:
        if isinstance(raw, str):
            model = ExtraStandardRequirement(requirement_name="", requirement_text=raw.strip(), source_section="")
        elif isinstance(raw, dict):
            model = ExtraStandardRequirement.model_validate(raw)
        else:
            continue
        key = (model.requirement_name, model.requirement_text, model.source_section)
        if key in seen or (not model.requirement_name and not model.requirement_text):
            continue
        seen.add(key)
        items.append(model)
    return items


def _normalize_string_list(value: Any) -> list[str]:
    if isinstance(value, list):
        raw_values = [str(item).strip() for item in value]
    elif isinstance(value, str):
        raw_values = [part.strip() for part in re.split(r"[,，;\n]+", value) if part.strip()]
    else:
        raw_values = []
    result: list[str] = []
    seen: set[str] = set()
    for item in raw_values:
        if item and item not in seen:
            seen.add(item)
            result.append(item)
    return result


def _normalize_quote_id(value: str, index: int, used_ids: set[str]) -> str:
    candidate = re.sub(r"[^a-zA-Z0-9_-]+", "_", value.strip()) or f"quote_{index}"
    candidate = candidate.strip("_-") or f"quote_{index}"
    if candidate not in used_ids:
        return candidate
    suffix = 2
    while f"{candidate}_{suffix}" in used_ids:
        suffix += 1
    return f"{candidate}_{suffix}"
