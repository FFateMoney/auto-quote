from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook


class ExcelPageRenderer:
    """Render an Excel workbook into page images for vision-based splitting."""

    def __init__(self, *, libreoffice_bin: str = "libreoffice", image_scale: float = 2.0, timeout_seconds: int = 120) -> None:
        self.libreoffice_bin = libreoffice_bin
        self.image_scale = image_scale
        self.timeout_seconds = timeout_seconds

    def render(self, workbook_path: Path, *, output_dir: Path) -> list[Path]:
        output_dir.mkdir(parents=True, exist_ok=True)
        pdf_path = self._convert_to_pdf(workbook_path, output_dir=output_dir)
        return self._render_pdf_pages(pdf_path, output_dir=output_dir)

    def _convert_to_pdf(self, workbook_path: Path, *, output_dir: Path) -> Path:
        with tempfile.TemporaryDirectory(prefix="auto_quote_excel_render_") as temp_name:
            temp_dir = Path(temp_name)
            prepared = temp_dir / workbook_path.name
            shutil.copy2(workbook_path, prepared)
            self._apply_print_layout(prepared)
            subprocess.run(
                [
                    self.libreoffice_bin,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(output_dir),
                    str(prepared),
                ],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                timeout=self.timeout_seconds,
            )
        pdf_path = output_dir / f"{workbook_path.stem}.pdf"
        if not pdf_path.exists():
            candidates = sorted(output_dir.glob("*.pdf"), key=lambda path: path.stat().st_mtime, reverse=True)
            if not candidates:
                raise RuntimeError(f"excel_pdf_render_failed:{workbook_path.name}")
            pdf_path = candidates[0]
        return pdf_path

    def _apply_print_layout(self, workbook_path: Path) -> None:
        workbook = load_workbook(workbook_path)
        for sheet in workbook.worksheets:
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.page_setup.orientation = "landscape"
        workbook.save(workbook_path)

    def _render_pdf_pages(self, pdf_path: Path, *, output_dir: Path) -> list[Path]:
        try:
            import pypdfium2 as pdfium
        except Exception as exc:
            raise RuntimeError("pypdfium2_not_available") from exc

        image_dir = output_dir / f"{pdf_path.stem}_pages"
        image_dir.mkdir(parents=True, exist_ok=True)
        pdf = pdfium.PdfDocument(str(pdf_path))
        try:
            paths: list[Path] = []
            for index in range(len(pdf)):
                page = pdf[index]
                bitmap = page.render(scale=self.image_scale)
                image = bitmap.to_pil()
                image_path = image_dir / f"page_{index + 1:03d}.png"
                image.save(image_path)
                paths.append(image_path)
            return paths
        finally:
            pdf.close()
