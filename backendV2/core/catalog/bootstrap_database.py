from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import psycopg
from openpyxl import load_workbook

from backendV2.core.catalog.fields import TEST_PROJECT_CAPABILITY_FIELDS
from backendV2.core.catalog.initialize_special_field_definitions import initialize_special_field_definitions
from backendV2.core.history.initialize import initialize_historical_quotations
from backendV2.core.settings import Settings


NONE_MARKERS = {None, "", "/", "none"}
MIGRATION_PATH = Path(__file__).parent / "migrations" / "001_replace_catalog.sql"
SOURCE_PATH = Path(__file__).resolve().parents[3] / "doc" / "新试验类型及设备能力.xlsx"
CAPABILITY_FIELD_NAMES = TEST_PROJECT_CAPABILITY_FIELDS


def rebuild_database(settings: Settings, source_path: Path = SOURCE_PATH) -> tuple[int, int]:
    workbook = load_workbook(source_path, data_only=True)
    project_rows = _read_test_projects(workbook["试验类型及设备匹配"])
    device_rows = _read_device_capabilities(workbook["设备能力"])
    _validate_device_references(project_rows, device_rows)
    capability_sets = _build_test_project_capability_sets(project_rows, device_rows)

    connection_args = {key: value for key, value in settings.database.items() if value not in (None, "")}
    with psycopg.connect(**connection_args) as connection:
        with connection.cursor() as cursor:
            cursor.execute(MIGRATION_PATH.read_text(encoding="utf-8"))
            cursor.executemany(
                """
                INSERT INTO public.test_projects (
                    standard_type, test_item, max_specification, pricing_mode,
                    base_fee, unit_price, applicable_device_codes
                ) VALUES (%(standard_type)s, %(test_item)s, %(max_specification)s, %(pricing_mode)s,
                          %(base_fee)s, %(unit_price)s, %(applicable_device_codes)s)
                """,
                project_rows,
            )
            cursor.executemany(
                """
                INSERT INTO public.device_capabilities (
                    device_code, max_volume_m3, max_length_mm, max_width_mm, max_height_mm,
                    temperature_min_c, temperature_max_c, humidity_min_rh, humidity_max_rh,
                    max_temperature_change_rate_c_per_min,
                    water_temperature_min_c, water_temperature_max_c,
                    water_flow_min_l_per_min, water_flow_max_l_per_min,
                    irradiance_min_w_per_m3, irradiance_max_w_per_m3,
                    max_load_kg, other_limits,
                    frequency_min_hz, frequency_max_hz,
                    acceleration_min_m_per_s2, acceleration_max_m_per_s2,
                    max_peak_to_peak_displacement_mm,
                    power_kwh
                ) VALUES (
                    %(device_code)s, %(max_volume_m3)s, %(max_length_mm)s, %(max_width_mm)s, %(max_height_mm)s,
                    %(temperature_min_c)s, %(temperature_max_c)s, %(humidity_min_rh)s, %(humidity_max_rh)s,
                    %(max_temperature_change_rate_c_per_min)s,
                    %(water_temperature_min_c)s, %(water_temperature_max_c)s,
                    %(water_flow_min_l_per_min)s, %(water_flow_max_l_per_min)s,
                    %(irradiance_min_w_per_m3)s, %(irradiance_max_w_per_m3)s,
                    %(max_load_kg)s, %(other_limits)s,
                    %(frequency_min_hz)s, %(frequency_max_hz)s,
                    %(acceleration_min_m_per_s2)s, %(acceleration_max_m_per_s2)s,
                    %(max_peak_to_peak_displacement_mm)s,
                    %(power_kwh)s
                )
                """,
                device_rows,
            )
            cursor.executemany(
                """
                INSERT INTO public.test_project_capability_sets (
                    test_project_id, capability_fields
                )
                SELECT id, %(capability_fields)s
                FROM public.test_projects
                WHERE standard_type = %(standard_type)s
                  AND test_item = %(test_item)s
                  AND max_specification = %(max_specification)s
                """,
                capability_sets,
            )
    initialize_special_field_definitions(settings)
    initialize_historical_quotations(settings)
    return len(project_rows), len(device_rows)


def _read_test_projects(sheet: Any) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in range(4, sheet.max_row + 1):
        devices = [
            str(sheet.cell(row, column).value).strip()
            for column in range(9, 21)
            if _value_or_none(sheet.cell(row, column).value) is not None
        ]
        rows.append(
            {
                "standard_type": _required_text(_merged_value(sheet, row, 3)),
                "test_item": _required_text(_merged_value(sheet, row, 4)),
                "max_specification": _normalize_max_specification(sheet.cell(row, 5).value),
                "pricing_mode": _required_text(sheet.cell(row, 6).value),
                "base_fee": sheet.cell(row, 7).value,
                "unit_price": sheet.cell(row, 8).value,
                "applicable_device_codes": devices,
            }
        )
    return rows


def _read_device_capabilities(sheet: Any) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in range(2, sheet.max_row + 1):
        raw_codes = _value_or_none(sheet.cell(row, 4).value)
        if raw_codes is None:
            continue
        values = {
            "max_volume_m3": _value_or_none(sheet.cell(row, 5).value),
            "max_length_mm": _value_or_none(sheet.cell(row, 6).value),
            "max_width_mm": _value_or_none(sheet.cell(row, 7).value),
            "max_height_mm": _value_or_none(sheet.cell(row, 8).value),
            "temperature_min_c": _range_value(sheet.cell(row, 9).value, 0),
            "temperature_max_c": _range_value(sheet.cell(row, 9).value, 1),
            "humidity_min_rh": _range_value(sheet.cell(row, 10).value, 0),
            "humidity_max_rh": _range_value(sheet.cell(row, 10).value, 1),
            "max_temperature_change_rate_c_per_min": _single_limit(sheet.cell(row, 11).value),
            "water_temperature_min_c": _range_value(sheet.cell(row, 12).value, 0),
            "water_temperature_max_c": _range_value(sheet.cell(row, 12).value, 1),
            "water_flow_min_l_per_min": _range_value(sheet.cell(row, 13).value, 0),
            "water_flow_max_l_per_min": _range_value(sheet.cell(row, 13).value, 1),
            "irradiance_min_w_per_m3": _range_value(sheet.cell(row, 14).value, 0),
            "irradiance_max_w_per_m3": _range_value(sheet.cell(row, 14).value, 1),
            "max_load_kg": _value_or_none(sheet.cell(row, 15).value),
            "other_limits": _value_or_none(sheet.cell(row, 16).value),
            "frequency_min_hz": _range_value(sheet.cell(row, 17).value, 0),
            "frequency_max_hz": _range_value(sheet.cell(row, 17).value, 1),
            "acceleration_min_m_per_s2": _range_value(sheet.cell(row, 18).value, 0),
            "acceleration_max_m_per_s2": _range_value(sheet.cell(row, 18).value, 1),
            "max_peak_to_peak_displacement_mm": _maximum_value(sheet.cell(row, 19).value),
            "power_kwh": _value_or_none(sheet.cell(row, 20).value),
        }
        for device_code in str(raw_codes).split("/"):
            rows.append({"device_code": device_code.strip(), **values})
    return rows


def _validate_device_references(project_rows: list[dict[str, Any]], device_rows: list[dict[str, Any]]) -> None:
    known_codes = {str(row["device_code"]) for row in device_rows}
    referenced_codes = {
        code
        for row in project_rows
        for code in row["applicable_device_codes"]
    }
    missing = sorted(referenced_codes - known_codes)
    if missing:
        raise ValueError(f"test projects reference unknown devices: {', '.join(missing)}")


def _build_test_project_capability_sets(
    project_rows: list[dict[str, Any]], device_rows: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    devices_by_code = {str(device["device_code"]): device for device in device_rows}
    capability_sets: list[dict[str, Any]] = []

    for project in project_rows:
        applicable_devices = [devices_by_code[code] for code in project["applicable_device_codes"]]
        capability_sets.append(
            {
                "standard_type": project["standard_type"],
                "test_item": project["test_item"],
                "max_specification": project["max_specification"],
                "capability_fields": [
                    field_name
                    for field_name in CAPABILITY_FIELD_NAMES
                    if any(device[field_name] is not None for device in applicable_devices)
                ],
            }
        )
    return capability_sets


def _merged_value(sheet: Any, row: int, column: int) -> Any:
    for merged_range in sheet.merged_cells.ranges:
        if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= column <= merged_range.max_col:
            return sheet.cell(merged_range.min_row, merged_range.min_col).value
    return sheet.cell(row, column).value


def _normalize_max_specification(value: Any) -> str:
    text = _required_text(value)
    matched = re.fullmatch(r"(?:≤)?\s*(\d+(?:\.\d+)?)\s*m3", text, flags=re.IGNORECASE)
    if not matched:
        return text
    numeric = float(matched.group(1))
    return str(int(numeric)) if numeric.is_integer() else str(numeric)


def _range_value(value: Any, index: int) -> float | None:
    text = _value_or_none(value)
    if text is None:
        return None
    matched = re.fullmatch(r"(-?\d+(?:\.\d+)?)\s*~\s*(-?\d+(?:\.\d+)?)", str(text))
    if not matched:
        if isinstance(text, (int, float)):
            return float(text)
        raise ValueError(f"expected a numeric range or value, got {text!r}")
    return float(matched.group(index + 1))


def _single_limit(value: Any) -> float | None:
    text = _value_or_none(value)
    if text is None:
        return None
    matched = re.fullmatch(r"≤\s*(-?\d+(?:\.\d+)?)", str(text))
    if not matched:
        raise ValueError(f"expected a single upper limit, got {text!r}")
    return float(matched.group(1))


def _maximum_value(value: Any) -> float | None:
    text = _value_or_none(value)
    if text is None:
        return None
    if isinstance(text, (int, float)):
        return float(text)
    matched = re.fullmatch(r"-?\d+(?:\.\d+)?\s*~\s*(-?\d+(?:\.\d+)?)", str(text))
    if not matched:
        raise ValueError(f"expected a numeric range or value, got {text!r}")
    return float(matched.group(1))


def _value_or_none(value: Any) -> Any | None:
    if value in NONE_MARKERS:
        return None
    return value


def _required_text(value: Any) -> str:
    normalized = _value_or_none(value)
    if normalized is None:
        return "none"
    return str(normalized).strip()


if __name__ == "__main__":
    project_count, device_count = rebuild_database(Settings.from_environment())
    print(f"rebuilt database: test_projects={project_count}, device_capabilities={device_count}")
