from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

from docx import Document


class QuotationDocumentExporter:
    """Writes the current quotation items into the maintained DOCX template."""

    def __init__(self, template_path: Path) -> None:
        self._template_path = template_path

    def export(self, output_path: Path, items: Sequence[dict[str, object]]) -> None:
        if not self._template_path.is_file():
            raise FileNotFoundError(f"Template not found: {self._template_path}")

        document = Document(self._template_path)
        if len(document.tables) > 1:
            self._fill_quote_table(document.tables[1], items)
        self._replace_placeholders(document, len(items))
        document.save(output_path)

    @staticmethod
    def _fill_quote_table(table: Any, items: Sequence[dict[str, object]]) -> None:
        start_row_index = next(
            (index for index, row in enumerate(table.rows) if row.cells and row.cells[0].text.strip() == "1"),
            -1,
        )
        if start_row_index < 0:
            return

        data_row_indices, total_row_index = _data_row_indices(table, start_row_index)
        if not data_row_indices:
            return
        template_row_index = data_row_indices[-1]
        for _ in range(max(0, len(items) - len(data_row_indices))):
            _insert_table_row_before(table, total_row_index if total_row_index >= 0 else len(table.rows), template_row_index)
        data_row_indices, total_row_index = _data_row_indices(table, start_row_index, include_empty=True)

        total_amount = 0.0
        for index, row_index in enumerate(data_row_indices):
            row = table.rows[row_index]
            if index >= len(items):
                for cell in row.cells:
                    cell.text = ""
                continue
            item = items[index]
            fixed_fields = item.get("fixed_fields")
            fields = fixed_fields if isinstance(fixed_fields, dict) else {}
            test_item = fields.get("raw_test_type") or ""
            row.cells[0].text = str(index + 1)
            row.cells[1].text = str(test_item)
            row.cells[2].text = ""
            row.cells[3].text = _number_text(item.get("base_fee"))
            row.cells[4].text = _number_text(item.get("unit_price"))
            row.cells[5].text = _number_text(item.get("pricing_quantity"))
            row.cells[6].text = _number_text(item.get("total_price"))
            total_amount += _number_value(item.get("total_price"))

        if total_row_index >= 0:
            table.rows[total_row_index].cells[6].text = _number_text(total_amount)

    @staticmethod
    def _replace_placeholders(document: Document, item_count: int) -> None:
        date_text = datetime.now().strftime("%Y年%m月%d日")
        for paragraph in document.paragraphs:
            if "20xx年x月x日" in paragraph.text:
                paragraph.text = paragraph.text.replace("20xx年x月x日", date_text)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if "xxxx" in cell.text or "x个测试项目" in cell.text:
                        cell.text = cell.text.replace("xxxx", "").replace("x个测试项目", f"{item_count}个测试项目")


def _data_row_indices(table: Any, start_row_index: int, *, include_empty: bool = False) -> tuple[list[int], int]:
    data_row_indices: list[int] = []
    total_row_index = -1
    for index in range(start_row_index, len(table.rows)):
        cell_text = table.rows[index].cells[0].text.strip()
        if cell_text.isdigit() or (include_empty and cell_text == ""):
            data_row_indices.append(index)
        elif "总计" in cell_text:
            total_row_index = index
            break
    return data_row_indices, total_row_index


def _insert_table_row_before(table: Any, row_index: int, template_row_index: int) -> None:
    template_row = table.rows[min(max(template_row_index, 0), len(table.rows) - 1)]
    new_row = deepcopy(template_row._tr)
    for cell in new_row.tc_lst:
        for paragraph in cell.p_lst:
            for run in paragraph.r_lst:
                for text in run.t_lst:
                    text.text = ""
    if row_index < len(table.rows):
        table.rows[max(row_index, 0)]._tr.addprevious(new_row)
    else:
        table._tbl.append(new_row)


def _number_value(value: object) -> float:
    try:
        return float(value) if value is not None else 0.0
    except (TypeError, ValueError):
        return 0.0


def _number_text(value: object) -> str:
    return f"{_number_value(value):g}"
