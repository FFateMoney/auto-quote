from __future__ import annotations

import json
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from docx import Document

from backendV2.core.catalog.repository import TestProject, TestProjectQuoteTemplate
from backendV2.core.catalog.special_fields import SpecialFieldDefinition
from backendV2.core.exports.quotation_docx import QuotationDocumentExporter
from backendV2.core.history.repository import HistoricalQuotation
from backendV2.core.quoting.service import QuotationResult
from backendV2.core.runs.agent_runner import AgentExecution
from backendV2.core.runs.service import RunService
from backendV2.core.settings import Settings
from backendV2.core.submissions.models import QuoteTable, QuoteTableSubmission
from backendV2.core.submissions.validator import QuoteTableValidator
from backendV2.core.tools.query_test_projects import QueryTestProjectsTool
from backendV2.core.tools.query_test_project_capability_fields import QueryTestProjectCapabilityFieldsTool
from backendV2.core.tools.query_test_project_aliases import QueryTestProjectAliasesTool
from backendV2.core.tools.submit_quote_batch import SubmitQuoteBatchTool
from backendV2.core.tools.validate_quote_table import ValidateQuoteTableTool


class InMemoryTestProjectRepository:
    def list_test_projects(self) -> list[TestProject]:
        return [
            TestProject(
                id=24,
                standard_type="振动",
                test_item="随机振动",
                max_specification="3吨台",
                pricing_mode="时长",
            )
        ]

    def get_quote_template(self, test_project_id: int) -> TestProjectQuoteTemplate:
        if test_project_id != 24:
            raise KeyError(test_project_id)
        return TestProjectQuoteTemplate(
            pricing_mode="时长",
            capability_fields=("frequency_min_hz", "frequency_max_hz", "max_peak_to_peak_displacement_mm"),
            max_specification="3吨台",
        )

    def get_test_project_aliases(self, test_project_id: int) -> tuple[str, ...]:
        if test_project_id != 24:
            raise KeyError(test_project_id)
        return ("振动测试", "Vibration Test")


class InMemorySpecialFieldRepository:
    def find_by_field_names(self, field_names: tuple[str, ...]) -> list[SpecialFieldDefinition]:
        comparison_types = {
            "frequency_min_hz": 0,
            "frequency_max_hz": 1,
            "max_peak_to_peak_displacement_mm": 1,
        }
        return [
            SpecialFieldDefinition(field_name, comparison_types[field_name])
            for field_name in field_names
            if field_name in comparison_types
        ]


class SuccessfulAgentRunner:
    def run(self, run_dir: Path, prompt_path: Path) -> AgentExecution:
        if not prompt_path.is_file():
            raise AssertionError("the Agent prompt must be available before execution")
        (run_dir / "submit_result.json").write_text(
            json.dumps({"accepted": True, "phase": "format", "message": "成功", "details": []}),
            encoding="utf-8",
        )
        return AgentExecution(return_code=0)


class RecordingRunService(RunService):
    def __init__(self, settings: Settings) -> None:
        super().__init__(settings, SuccessfulAgentRunner())
        self.received_batches: list[object] = []

    def accept_batch_submission(self, run_dir: Path, quote_tables: list[object]) -> None:
        del run_dir
        self.received_batches.append(quote_tables)


class RecordingQuotationProcessor:
    def quote_batch(self, quote_tables: list[QuoteTableSubmission]) -> list[QuotationResult]:
        return [
            QuotationResult(
                quote_id=table.quote_id,
                test_project_id=table.quote_table.test_project_id,
                eligible_device_codes=("F10",),
                selected_device_code="F10",
                base_fee=800,
                unit_price=200,
                pricing_quantity=15,
                amount=3800,
            )
            for table in quote_tables
        ]

    def quote_values(
        self,
        quote_id: str,
        test_project_id: int | None,
        fixed_fields: dict[str, object],
        special_fields: dict[str, object],
        selected_device_code: str | None = None,
    ) -> QuotationResult:
        del special_fields
        quantity = fixed_fields.get("pricing_quantity")
        amount = 800 + 200 * float(quantity) if isinstance(quantity, (int, float)) else None
        return QuotationResult(
            quote_id=quote_id,
            test_project_id=test_project_id,
            eligible_device_codes=("F10",),
            selected_device_code=selected_device_code or "F10",
            base_fee=800,
            unit_price=200,
            pricing_quantity=quantity,
            amount=amount,
            pricing_mode="时长",
            specification="3吨台",
            standard_type="振动",
            test_item="振动",
        )


class RecordingDocumentExporter:
    def __init__(self) -> None:
        self.exports: list[list[dict[str, object]]] = []

    def export(self, output_path: Path, items: list[dict[str, object]]) -> None:
        self.exports.append(items)
        output_path.write_bytes(b"quotation document")


class RecordingAliasWriter:
    def __init__(self) -> None:
        self.calls: list[dict[int, tuple[str, ...]]] = []

    def append_test_project_aliases(self, aliases_by_project: dict[int, tuple[str, ...]]) -> None:
        self.calls.append(aliases_by_project)


class RecordingHistoryWriter:
    def __init__(self) -> None:
        self.saved_quotations: list[HistoricalQuotation] = []

    def save_historical_quotations(self, quotations: list[HistoricalQuotation]) -> int:
        self.saved_quotations = list(quotations)
        return len(quotations)


class RunToolsTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        root = Path(self.temporary_directory.name)
        self.settings = Settings(
            project_root=root,
            core_runtime_root=root / "core" / "runtime",
            agent_workspace_root=root / "agent_workspace",
            agent_executable="codex",
            agent_profile="quote-agent",
            qwen_api_key="",
            database={},
        )
        self.settings.agent_prompts_root.mkdir(parents=True)
        (self.settings.agent_prompts_root / "extract_quote.md").write_text("test prompt", encoding="utf-8")
        self.validator = QuoteTableValidator()

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_query_validate_and_submit_a_batch_in_one_run_workspace(self) -> None:
        run_dir = self._create_run_dir()
        first_table_path = run_dir / "quotes" / "quote-001" / "quote_table.json"
        second_table_path = run_dir / "quotes" / "quote-002" / "quote_table.json"
        first_table_path.parent.mkdir(parents=True)
        second_table_path.parent.mkdir(parents=True)
        first_table_path.write_text(json.dumps(self._valid_table(10), ensure_ascii=False), encoding="utf-8")
        second_table_path.write_text(json.dumps(self._valid_table(None), ensure_ascii=False), encoding="utf-8")

        catalog_path = QueryTestProjectsTool(self.settings, InMemoryTestProjectRepository()).execute(run_dir)
        capability_path = QueryTestProjectCapabilityFieldsTool(
            self.settings,
            InMemoryTestProjectRepository(),
            InMemorySpecialFieldRepository(),
        ).execute(run_dir, 24)
        alias_path = QueryTestProjectAliasesTool(
            self.settings,
            InMemoryTestProjectRepository(),
        ).execute(run_dir, 24)
        validation, validation_path = ValidateQuoteTableTool(self.settings, self.validator).execute(
            run_dir, "quotes/quote-001/quote_table.json"
        )
        runs = RecordingRunService(self.settings)
        submission, submission_path = SubmitQuoteBatchTool(
            self.settings,
            self.validator,
            runs,
        ).execute(run_dir)

        self.assertEqual(
            json.loads(catalog_path.read_text(encoding="utf-8"))["items"][0],
            {
                "id": 24,
                "standard_type": "振动",
                "test_item": "随机振动",
                "max_specification": "3吨台",
                "pricing_mode": "时长",
            },
        )
        self.assertEqual(
            json.loads(capability_path.read_text(encoding="utf-8")),
            {
                "固定字段": {
                    "raw_test_type": None,
                    "standard_code": None,
                    "standard_document_section": None,
                    "pricing_mode": "时长",
                    "specification": "3吨台",
                    "pricing_quantity": None,
                    "sample_count": None,
                    "length_mm": None,
                    "width_mm": None,
                    "height_mm": None,
                },
                "专有字段": {
                    "frequency_min_hz": None,
                    "frequency_max_hz": None,
                    "max_peak_to_peak_displacement_mm": None,
                },
            },
        )
        self.assertEqual(
            json.loads(alias_path.read_text(encoding="utf-8")),
            {
                "test_project_id": 24,
                "aliases": ["振动测试", "Vibration Test"],
            },
        )
        self.assertTrue(validation.accepted)
        self.assertTrue(submission.accepted)
        self.assertEqual(json.loads(validation_path.read_text(encoding="utf-8"))["message"], "成功")
        self.assertEqual(json.loads(submission_path.read_text(encoding="utf-8"))["message"], "成功")
        self.assertEqual(validation_path, first_table_path.with_name("validation_result.json"))
        self.assertEqual([[table.quote_id for table in batch] for batch in runs.received_batches], [["quote-001", "quote-002"]])

    def test_submit_batch_reports_invalid_quote_table_format(self) -> None:
        run_dir = self._create_run_dir()
        table_path = run_dir / "quotes" / "quote-001" / "quote_table.json"
        table_path.parent.mkdir(parents=True)
        table_path.write_text('{"schema_version": 1, "values": {}}', encoding="utf-8")

        validation, _ = ValidateQuoteTableTool(self.settings, self.validator).execute(run_dir, "quotes/quote-001/quote_table.json")
        submission, submission_path = SubmitQuoteBatchTool(
            self.settings,
            self.validator,
            RunService(self.settings, SuccessfulAgentRunner()),
        ).execute(run_dir)

        self.assertFalse(validation.accepted)
        self.assertFalse(submission.accepted)
        self.assertEqual(
            [detail["rule"] for detail in validation.details],
            [detail["rule"] for detail in submission.details],
        )
        self.assertTrue(all(detail["path"].startswith("quotes/quote-001/quote_table.json:") for detail in submission.details))
        self.assertEqual(json.loads(submission_path.read_text(encoding="utf-8"))["message"], "格式错误")

    def test_validate_rejects_a_directory_outside_agent_workspace(self) -> None:
        outside = Path(self.temporary_directory.name) / "outside"
        outside.mkdir()
        with self.assertRaisesRegex(ValueError, "agent_workspace/runtime"):
            ValidateQuoteTableTool(self.settings, self.validator).execute(outside, "quote_table.json")

    def test_run_service_creates_artifacts_and_records_agent_completion(self) -> None:
        service = RunService(self.settings, SuccessfulAgentRunner())
        created = service.create_run("客户报价.xlsx", b"placeholder Excel bytes")
        run_dir = self.settings.agent_runs_root / created.run_id

        completed = service.run_agent(created.run_id)

        self.assertTrue((run_dir / "input" / "客户报价.xlsx").is_file())
        self.assertTrue((run_dir / "quote_table.schema.json").is_file())
        self.assertEqual(completed.status, "submitted")
        self.assertEqual(completed.agent_return_code, 0)

    def test_run_service_writes_core_quotation_results_after_submission(self) -> None:
        run_dir = self._create_run_dir()
        service = RunService(self.settings, SuccessfulAgentRunner(), RecordingQuotationProcessor())
        table = QuoteTable.model_validate(self._valid_table(24))

        service.accept_batch_submission(run_dir, [QuoteTableSubmission("quote-001", table)])

        self.assertEqual(
            json.loads((run_dir / "quotation_results.json").read_text(encoding="utf-8")),
            {
                "items": [
                    {
                        "quote_id": "quote-001",
                        "test_project_id": 24,
                        "eligible_device_codes": ["F10"],
                        "selected_device_code": "F10",
                        "base_fee": 800.0,
                        "unit_price": 200.0,
                        "pricing_quantity": 15,
                        "amount": 3800.0,
                        "pricing_mode": None,
                        "specification": None,
                        "standard_type": None,
                        "test_item": None,
                        "specification_options": [],
                        "eligible_devices": [],
                    }
                ]
            },
        )

    def test_manual_changes_create_snapshots_and_export_writes_the_current_output_snapshot(self) -> None:
        run_dir = self._create_run_dir()
        quote_dir = run_dir / "quotes" / "quote-001"
        quote_dir.mkdir(parents=True)
        table_payload = self._valid_table(24)
        table_payload["values"]["固定字段"] = {
            "raw_test_type": "Vibration Testing",
            "pricing_quantity": 5,
        }
        quote_path = quote_dir / "quote_table.json"
        quote_path.write_text(json.dumps(table_payload, ensure_ascii=False), encoding="utf-8")
        table = QuoteTable.model_validate(table_payload)
        exporter = RecordingDocumentExporter()
        aliases = RecordingAliasWriter()
        service = RunService(
            self.settings,
            SuccessfulAgentRunner(),
            RecordingQuotationProcessor(),
            exporter,
            aliases,
        )
        service.accept_batch_submission(run_dir, [QuoteTableSubmission("quote-001", table)])

        first = service.save_quotation_snapshot(
            "test-run",
            "quote-001",
            24,
            {"raw_test_type": "振动测试", "pricing_quantity": 6},
            {"frequency_max_hz": 2000},
        )
        second = service.save_quotation_snapshot(
            "test-run",
            "quote-001",
            24,
            {"raw_test_type": "振动测试", "pricing_quantity": 7},
            {"frequency_max_hz": 2100},
            "F10",
        )

        snapshots = sorted((quote_dir / "snapshots").glob("quotation_snapshot_*.json"))
        self.assertEqual(len(snapshots), 2)
        self.assertNotEqual(json.loads(snapshots[0].read_text(encoding="utf-8"))["id"], json.loads(snapshots[1].read_text(encoding="utf-8"))["id"])
        self.assertEqual(first["pricing_quantity"], 6)
        self.assertEqual(second["pricing_quantity"], 7)
        self.assertEqual(service.get_quotation_results("test-run")["items"][0]["pricing_quantity"], 7)

        first_export = service.export_quotation_document("test-run")
        first_output_snapshot = json.loads((run_dir / "quotation_results_output.json").read_text(encoding="utf-8"))
        second_export = service.export_quotation_document("test-run", quote_id="quote-001")
        second_output_snapshot = json.loads((run_dir / "quotation_results_output.json").read_text(encoding="utf-8"))

        self.assertTrue(first_export.is_file())
        self.assertTrue(second_export.is_file())
        self.assertEqual(len(exporter.exports), 2)
        self.assertEqual(first_output_snapshot["quote_id"], None)
        self.assertEqual(second_output_snapshot["quote_id"], "quote-001")
        self.assertNotEqual(first_output_snapshot["id"], second_output_snapshot["id"])
        self.assertEqual(aliases.calls, [{24: ("振动测试",)}, {24: ("振动测试",)}])

    def test_saves_the_latest_manual_quotation_snapshot_to_history(self) -> None:
        history_writer = RecordingHistoryWriter()
        service = RunService(
            self.settings,
            SuccessfulAgentRunner(),
            RecordingQuotationProcessor(),
            history_writer=history_writer,
        )
        created = service.create_run("客户报价.xlsx", b"requirement")
        run_dir = self.settings.agent_runs_root / created.run_id
        quote_dir = run_dir / "quotes" / "quote-001"
        quote_dir.mkdir(parents=True)
        table_payload = self._valid_table(24)
        table_payload["values"]["固定字段"] = {
            "raw_test_type": "振动测试",
            "standard_code": "MIL-STD 810G",
            "pricing_quantity": 5,
            "sample_count": 1,
            "length_mm": 100,
            "width_mm": 80,
            "height_mm": 20,
        }
        (quote_dir / "quote_table.json").write_text(json.dumps(table_payload, ensure_ascii=False), encoding="utf-8")
        service.accept_batch_submission(run_dir, [QuoteTableSubmission("quote-001", QuoteTable.model_validate(table_payload))])
        service.save_quotation_snapshot(
            created.run_id,
            "quote-001",
            24,
            {
                "raw_test_type": "人工修改后的振动测试",
                "standard_code": "MIL-STD 810H",
                "pricing_quantity": 6,
                "sample_count": 2,
                "length_mm": 120,
                "width_mm": 90,
                "height_mm": 30,
            },
            {"frequency_max_hz": 500},
        )

        saved_count = service.save_current_quotations_to_history(created.run_id)

        self.assertEqual(saved_count, 1)
        self.assertEqual(len(history_writer.saved_quotations), 1)
        saved = history_writer.saved_quotations[0]
        self.assertEqual(saved.quotation_run_id, created.run_id)
        self.assertEqual(saved.quote_id, "quote-001")
        self.assertIsNotNone(saved.source_snapshot_id)
        self.assertEqual(saved.source_file_name, "客户报价.xlsx")
        self.assertEqual(saved.raw_test_type, "人工修改后的振动测试")
        self.assertEqual(saved.standard_code, "MIL-STD 810H")
        self.assertEqual(saved.pricing_quantity, 6)
        self.assertEqual(saved.special_fields, {"frequency_max_hz": 500})
        self.assertEqual(saved.quotation_snapshot["fixed_fields"]["raw_test_type"], "人工修改后的振动测试")

    def test_docx_exporter_adds_rows_before_the_total(self) -> None:
        template_path = Path(self.temporary_directory.name) / "quote_template.docx"
        document = Document()
        document.add_table(rows=1, cols=1)
        table = document.add_table(rows=3, cols=7)
        table.rows[0].cells[0].text = "序号"
        table.rows[1].cells[0].text = "1"
        table.rows[2].cells[0].text = "总计"
        document.save(template_path)
        output_path = Path(self.temporary_directory.name) / "报价单.docx"

        QuotationDocumentExporter(template_path).export(
            output_path,
            [
                {
                    "fixed_fields": {
                        "raw_test_type": f"原始类型 {index}",
                        "test_item": f"标准测试项目 {index}",
                    },
                    "base_fee": 100,
                    "unit_price": 20,
                    "pricing_quantity": index,
                    "total_price": 100 + 20 * index,
                }
                for index in range(1, 4)
            ],
        )

        exported_table = Document(output_path).tables[1]
        self.assertEqual([exported_table.rows[index].cells[0].text for index in range(1, 4)], ["1", "2", "3"])
        self.assertEqual([exported_table.rows[index].cells[1].text for index in range(1, 4)], ["原始类型 1", "原始类型 2", "原始类型 3"])
        self.assertEqual(exported_table.rows[4].cells[0].text, "总计")
        self.assertEqual(exported_table.rows[4].cells[6].text, "420")

    def test_base_fee_override_can_be_enabled_and_restored_for_all_quotations(self) -> None:
        run_dir = self._create_run_dir()
        quote_dir = run_dir / "quotes" / "quote-001"
        quote_dir.mkdir(parents=True)
        (quote_dir / "quote_table.json").write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "test_project_id": 24,
                    "values": {
                        "固定字段": {"raw_test_type": "振动", "pricing_quantity": 15},
                        "专有字段": {"frequency_max_hz": 2000},
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        (run_dir / "quotation_results.json").write_text(
            json.dumps({"items": [{"quote_id": "quote-001"}]}, ensure_ascii=False),
            encoding="utf-8",
        )
        service = RunService(self.settings, quotation_processor=RecordingQuotationProcessor())
        service.save_quotation_snapshot(
            "test-run",
            "quote-001",
            24,
            {"raw_test_type": "振动", "pricing_quantity": 15},
            {"frequency_max_hz": 2000},
        )

        results = service.set_all_quotation_base_fee_overrides("test-run", Decimal("0"))

        self.assertEqual(results["items"][0]["base_fee"], 0.0)
        self.assertEqual(results["items"][0]["base_fee_override"], 0.0)
        self.assertEqual(results["items"][0]["total_price"], 3000.0)

        restored = service.set_all_quotation_base_fee_overrides("test-run", None)

        self.assertEqual(restored["items"][0]["base_fee"], 800)
        self.assertNotIn("base_fee_override", restored["items"][0])
        self.assertEqual(restored["items"][0]["total_price"], 3800.0)

    def _create_run_dir(self) -> Path:
        run_dir = self.settings.agent_runs_root / "test-run"
        run_dir.mkdir(parents=True)
        return run_dir

    @staticmethod
    def _valid_table(test_project_id: int | None) -> dict[str, object]:
        return {
            "schema_version": 1,
            "test_project_id": test_project_id,
            "values": {
                "固定字段": {
                    "length_mm": 1143,
                    "width_mm": 488.95,
                    "height_mm": 209.55,
                },
                "专有字段": {},
            },
        }


if __name__ == "__main__":
    unittest.main()
